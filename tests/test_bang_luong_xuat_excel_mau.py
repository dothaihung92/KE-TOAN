import os
import sys
import tempfile

_REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, _REPO_ROOT)
import server
import openpyxl

# Yêu cầu: "thêm xuất Excel theo mẫu" — file Excel theo mẫu sheet BL của người dùng: mỗi tháng 1 sheet, khối trái BẢNG TÍNH
# LƯƠNG, khối phải BẢNG CHẤM CÔNG (X/L, TNC), có công thức, đặt sẵn khổ giấy + vùng in + lặp tiêu đề.

TS = server._luong_chuan_tham_so(None, 2024)


def dong(ma, ten, **kw):
    d = {"ma": ma, "ten": ten, "luong_cb": 5_310_000, "tien_com": 700_000, "muc_xang": 500_000, "pc_chuc_vu": 500_000, "muc_dt": 500_000,
         "trang_phuc": 400_000}
    d.update(kw)
    return d


rows_nhap = [dong("2", "Nguyễn Văn A", dong_bh=1), dong("3", "Lê Thị B", dong_bh=1, ngay_lam=20),
             dong("4", "Trần C", dong_bh=0, ngay_lam=5, thuong_bh=1_500_000)]
tinh = [server._luong_tinh_dong(r, TS, "05") for r in rows_nhap]

# dựng payload đúng như giao diện (cột hiển thị + chấm công) — cột toàn 0 bị ẩn
ky_hieu = [("luong_cb", "Lương căn bản"), ("ngay_cong_hd", "Ngày công"), ("ngay_lam_hd", "Tổng NC"), ("luong", "Tiền lương"),
           ("tt_tien_com", "Tiền cơm"), ("xang_xe", "Xăng xe"), ("dien_thoai", "Điện thoại"),
           ("tt_trang_phuc", "Trang phục"), ("thuong_bh", "Thưởng bán hàng"), ("bhxh_nld", "BHXH 8%"), ("bhyt_nld", "BHYT 1.5%"),
           ("bhtn_nld", "BHTN 1%"), ("thue_tru_luong", "Thuế TNCN"), ("tt_luong", "Tổng thực nhận")]
nhom = {"tt_tien_com": "pc", "xang_xe": "pc", "dien_thoai": "pc", "tt_trang_phuc": "pc", "thuong_bh": "pc",
        "bhxh_nld": "gt", "bhyt_nld": "gt", "bhtn_nld": "gt", "thue_tru_luong": "gt"}
cot = [{"k": "stt", "t": "Stt", "w": 26}, {"k": "ma", "t": "Mã NV", "w": 40}, {"k": "ten", "t": "Họ và Tên", "w": 150}, {"k": "chuc_vu", "t": "Chức vụ", "w": 88}]
for k, t in ky_hieu:
    cot.append({"k": k, "t": t, "nhom": nhom.get(k, ""), "n": k not in ("ngay_cong_hd", "ngay_lam_hd"), "dp": k in ("ngay_cong_hd", "ngay_lam_hd"),
                "tong": k not in ("ngay_cong_hd", "ngay_lam_hd"), "dam": k == "tt_luong", "w": 80})
cot.append({"k": "ky", "t": "Ký nhận", "w": 96})
dim = 31
ngay = []
import datetime
for dd in range(1, dim + 1):
    thu = (datetime.date(2024, 5, dd).weekday() + 1) % 7          # 0 = CN
    ngay.append({"d": dd, "thu": thu, "cn": thu == 0, "le": dd == 1})


def cham_cua(k):
    g = tinh[k]["ngay_lam_hd"]
    dau, dem = [], 0
    for x in ngay:
        if x["cn"]:
            dau.append("")
        elif x["le"]:
            dau.append("L")
        elif dem < g:
            dau.append("X")
            dem += 1
        else:
            dau.append("")
    return {"tong": g, "soX": dem, "dau": dau}


thang_payload = {"05": {"cot": cot,
                        "rows": [dict({"ten": r["ten"], "ma": r["ma"], "chuc_vu": "KD", "thoi_vu": t["thoi_vu"]},
                                      **{c["k"]: t[c["k"]] for c in cot if c["k"] not in ("stt", "ky", "ma", "ten", "chuc_vu")}) for r, t in zip(rows_nhap, tinh)],
                        "cham": [cham_cua(k) for k in range(3)], "ngay": ngay}}
thang_payload["06"] = {"cot": cot, "rows": thang_payload["05"]["rows"][:2], "cham": [cham_cua(0), cham_cua(1)], "ngay": ngay[:30]}
tuy = {"ten_cty": "CÔNG TY TNHH A", "dia_chi": "1/50 Thanh Đa", "mst": "0312345678", "nguoi_lap": "Đỗ Thái Hưng", "giam_doc": "Nguyễn Văn Kiên",
       "in_luong": True, "in_cong": True, "kho": "a4n"}

_dl = server.DOWNLOAD_DIR
server.DOWNLOAD_DIR = tempfile.mkdtemp()
try:
    duong, ten = server._luong_xuat_excel_mau(2024, thang_payload, tuy)
    assert ten == "BangLuong_ChamCong_2024.xlsx" and os.path.exists(duong)
    wb = openpyxl.load_workbook(duong)
    assert wb.sheetnames == ["BL 05-2024", "BL 06-2024"], wb.sheetnames
    ws = wb["BL 05-2024"]
    nL = len(cot)

    # ===== 1: khối lương: đầu trang, tiêu đề, tiêu đề cột 2 tầng, dữ liệu, tổng cộng, chữ ký =====
    assert ws["A1"].value == "CÔNG TY TNHH A" and ws["A2"].value == "ĐC: 1/50 Thanh Đa" and ws["A3"].value == "MST: 0312345678"
    assert ws["A4"].value == "BẢNG TÍNH LƯƠNG VÀ CÁC KHOẢN THU NHẬP KHÁC" and ws["A5"].value == "THÁNG 05 NĂM 2024"
    tieu_de_8 = {ws.cell(8, j).value for j in range(1, nL + 1)}
    assert {"Stt", "Họ và Tên", "Lương căn bản", "Phụ cấp", "Các khoản giảm trừ", "Tổng thực nhận", "Ký nhận"} <= tieu_de_8, tieu_de_8
    assert {ws.cell(9, j).value for j in range(1, nL + 1)} >= {"Tiền cơm", "BHXH 8%", "Thuế TNCN"}
    assert ws["C10"].value == "Nguyễn Văn A" and ws["A10"].value == 1 and ws["A12"].value == 3
    assert ws["A13"].value == "Tổng cộng" and any(str(ws.cell(13, j).value or "").startswith("=SUM(") for j in range(1, nL + 1))
    txt = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    assert "Người lập biểu" in txt and "Giám đốc" in txt and "Đỗ Thái Hưng" in txt and "Nguyễn Văn Kiên" in txt and "Ngày 31 tháng 05 năm 2024" in txt
    print("PASS 1: sheet BL 05-2024 có đầu trang, tiêu đề, cột 2 tầng (Phụ cấp / Các khoản giảm trừ), dữ liệu, Tổng cộng, chữ ký.")

    # ===== 2: khối chấm công liền sau khối lương: X/L, CN tô xám, TNC = công thức COUNTIF, thời vụ ghi chú =====
    c1 = nL + 1
    assert ws.cell(4, c1).value == "BẢNG CHẤM CÔNG THÁNG TỪ 01/05/2024  ĐẾN 31/05/2024"
    assert ws.cell(8, c1).value == "STT" and ws.cell(8, c1 + 1).value == "Họ và Tên"
    assert ws.cell(8, c1 + 2).value == "T4" and ws.cell(9, c1 + 2).value == 1, "Ngày 1/5/2024 là thứ Tư"
    assert ws.cell(9, c1 + 2 + 30).value == 31 and ws.cell(8, c1 + 2 + 4).value == "CN" and ws.cell(9, c1 + 2 + 4).value == 5
    assert ws.cell(10, c1 + 2).value == "L", "Ngày lễ"
    assert ws.cell(10, c1 + 2 + 4).value is None and ws.cell(10, c1 + 2 + 4).fill.fgColor.rgb.endswith("D9D9D9"), "Chủ nhật trống + tô xám"
    x_dem = sum(1 for j in range(31) if ws.cell(10, c1 + 2 + j).value == "X")
    assert x_dem == tinh[0]["ngay_lam_hd"] == 26
    c_tnc = c1 + 2 + 31
    assert str(ws.cell(10, c_tnc).value).startswith("=COUNTIF(") and ws.cell(8, c_tnc).value == "TNC" and ws.cell(8, c_tnc + 1).value == "Ghi chú"
    assert ws.cell(12, c_tnc + 1).value == "Thời vụ (không BHXH)" and ws.cell(10, c_tnc + 1).value is None
    assert ws.cell(4, c1 + 2).value is None
    print("PASS 2: khối chấm công: thứ + ngày, X/L, Chủ nhật xám, TNC = COUNTIF, ghi chú thời vụ.")

    # ===== 3: tháng 6 (30 ngày, 2 người) đúng số cột ngày + dòng tổng đúng chỗ =====
    w6 = wb["BL 06-2024"]
    assert w6.cell(9, c1 + 2 + 29).value == 30 and w6.cell(8, c1 + 2 + 30).value == "TNC", "Tháng 30 ngày: cột TNC ngay sau ngày 30"
    assert w6["A12"].value == "Tổng cộng" and w6["A10"].value == 1 and w6["A11"].value == 2
    print("PASS 3: tháng 30 ngày có đúng 30 cột ngày; dòng Tổng cộng ngay dưới người cuối.")

    # ===== 4: tính lại công thức Excel độc lập: Tổng thực nhận, Tổng cộng, TNC khớp phần mềm =====
    try:
        import formulas
    except ImportError:
        formulas = None
    if formulas:
        sol = formulas.ExcelModel().loads(duong).finish().calculate()
        gt = lambda sheet, ref: float(list(sol["'[%s]%s'!%s" % (os.path.basename(duong), sheet.upper(), ref)].value[0])[0])
        pos = {c["k"]: j for j, c in enumerate(cot, 1)}
        L = openpyxl.utils.get_column_letter
        for i_r, t in enumerate(tinh):
            assert round(gt("BL 05-2024", f"{L(pos['tt_luong'])}{10 + i_r}")) == t["tt_luong"], (i_r, t["tt_luong"])
        tong_tt = sum(t["tt_luong"] for t in tinh)
        assert round(gt("BL 05-2024", f"{L(pos['tt_luong'])}13")) == tong_tt
        assert round(gt("BL 05-2024", f"{L(pos['luong'])}13")) == round(sum(t["luong"] for t in tinh))
        assert gt("BL 05-2024", f"{L(c_tnc)}10") == 26 and gt("BL 05-2024", f"{L(c_tnc)}11") == 20 and gt("BL 05-2024", f"{L(c_tnc)}12") == 5
        assert tinh[2]["thue_tru_luong"] > 0 and tinh[2]["tt_luong"] == tinh[2]["chi_phi_luong"] - tinh[2]["thue_tru_luong"], "Người thời vụ bị trừ thuế 10%"
        print("PASS 4: công thức Excel tính lại độc lập khớp phần mềm (thực nhận từng người, tổng cộng, TNC).")

    # ===== 5: thiết lập in: khổ giấy, hướng, vừa bề ngang, vùng in 2 khối (2 trang), lặp tiêu đề =====
    assert ws.page_setup.orientation == "landscape" and ws.page_setup.paperSize == 9
    assert ws.sheet_properties.pageSetUpPr.fitToPage is True and ws.page_setup.fitToWidth == 1 and ws.page_setup.fitToHeight == 0
    assert str(ws.print_title_rows).replace("$", "") == "8:9", ws.print_title_rows
    vung = str(ws.print_area)
    assert vung.count(":") == 2 and "$A$1:$" in vung and f"${L(c1)}$1:$" in vung, vung
    print("PASS 5: A4 ngang, vừa bề ngang, 2 vùng in (lương / chấm công), lặp tiêu đề hàng 8-9.")
    for kho, huong, giay in (("a3n", "landscape", 8), ("a4d", "portrait", 9)):
        w2 = openpyxl.load_workbook(server._luong_xuat_excel_mau(2024, thang_payload, dict(tuy, kho=kho))[0])["BL 05-2024"]
        assert (w2.page_setup.orientation, w2.page_setup.paperSize) == (huong, giay), kho

    # ===== 6: chỉ bảng lương / chỉ bảng chấm công =====
    wl = openpyxl.load_workbook(server._luong_xuat_excel_mau(2024, thang_payload, dict(tuy, in_cong=False))[0])["BL 05-2024"]
    assert wl.max_column == nL and str(wl.print_area).count(":") == 1
    wc = openpyxl.load_workbook(server._luong_xuat_excel_mau(2024, thang_payload, dict(tuy, in_luong=False))[0])["BL 05-2024"]
    assert wc["A4"].value.startswith("BẢNG CHẤM CÔNG") and wc["A8"].value == "STT" and wc["B10"].value == "Nguyễn Văn A"
    # ===== 7: tháng không có dữ liệu bị bỏ qua; không có tháng nào -> lỗi 404 =====
    from fastapi import HTTPException
    w3 = openpyxl.load_workbook(server._luong_xuat_excel_mau(2024, {"05": thang_payload["05"], "06": {"cot": cot, "rows": [], "cham": [], "ngay": ngay}}, tuy)[0])
    assert w3.sheetnames == ["BL 05-2024"]
    try:
        server._luong_xuat_excel_mau(2024, {}, tuy)
        raise AssertionError("phải báo lỗi")
    except HTTPException as e:
        assert e.status_code == 404
    print("PASS 6-7: chỉ lương / chỉ chấm công; tháng trống bị bỏ qua; không có dữ liệu -> báo lỗi.")
finally:
    server.DOWNLOAD_DIR = _dl

# ===== 8: bảng chấm công có GIỜ TĂNG CA: ô "X+n", cột Giờ TC, TNC vẫn đếm đúng (COUNTIF "X*"); số ngày/giờ hiện dạng số =====
import copy
p2 = copy.deepcopy(thang_payload)
tc0 = [0.0] * 31
for d_, h_ in ((2, 2.0), (7, 1.5), (9, 4.0), (14, 4.4)):
    tc0[d_ - 1] = h_
p2["05"]["cham"][0]["tc"] = tc0
p2["05"]["cham"][0]["gio"] = 11.9
_dl2 = server.DOWNLOAD_DIR
server.DOWNLOAD_DIR = tempfile.mkdtemp()
try:
    duong2, _ = server._luong_xuat_excel_mau(2024, p2, tuy)
    w = openpyxl.load_workbook(duong2)["BL 05-2024"]
    nL2 = len(cot)
    c1 = nL2 + 1
    assert w.cell(10, c1 + 2 + 1).value == "X+2" and w.cell(10, c1 + 2 + 6).value == "X+1,5" and w.cell(10, c1 + 2 + 13).value == "X+4,4"
    assert w.cell(10, c1 + 2 + 2).value == "X", "Ngày không tăng ca vẫn là X"
    c_tnc = c1 + 2 + 31
    assert w.cell(8, c_tnc).value == "TNC" and w.cell(8, c_tnc + 1).value == "Giờ TC" and w.cell(8, c_tnc + 2).value == "Ghi chú"
    assert w.cell(10, c_tnc + 1).value == 11.9 and w.cell(11, c_tnc + 1).value is None, "Chỉ người có tăng ca mới có giờ"
    assert str(w.cell(10, c_tnc).value).startswith('=COUNTIF(') and str(w.cell(10, c_tnc).value).endswith(',"X*")')
    assert "X+n" in str(w.cell(14, c1).value)
    assert w.column_dimensions[openpyxl.utils.get_column_letter(c1 + 2)].width > 5, "Cột ngày rộng hơn để vừa 'X+1,5'"
    # định dạng số: "24" không còn "24."; giờ lẻ hiện 1 số thập phân
    pos = {c["k"]: j for j, c in enumerate(cot, 1)}
    assert w.cell(10, pos["ngay_cong_hd"]).number_format == "0" and w.cell(10, pos["ngay_lam_hd"]).number_format == "0"
    assert w.cell(11, pos["ngay_lam_hd"]).value == 20 and w.cell(11, pos["ngay_lam_hd"]).number_format == "0"
    assert w.cell(10, c_tnc).number_format == "General" and w.cell(10, c_tnc + 1).number_format == "General"
    assert w.column_dimensions["C"].width >= 28, "Cột họ tên đủ rộng, không bị cắt"
    try:
        import formulas
        sol = formulas.ExcelModel().loads(duong2).finish().calculate()
        ref = "'[%s]%s'!%s10" % (os.path.basename(duong2), "BL 05-2024".upper(), openpyxl.utils.get_column_letter(c_tnc))
        assert float(list(sol[ref].value[0])[0]) == 26, "TNC đếm cả ô X+n"
    except ImportError:
        pass
    # payload không có giờ tăng ca -> không thêm cột Giờ TC, cột ngày hẹp như cũ
    w0 = openpyxl.load_workbook(server._luong_xuat_excel_mau(2024, thang_payload, tuy)[0])["BL 05-2024"]
    assert w0.cell(8, c_tnc + 1).value == "Ghi chú" and w0.column_dimensions[openpyxl.utils.get_column_letter(c1 + 2)].width < 5
finally:
    server.DOWNLOAD_DIR = _dl2
print("PASS 8: chấm công có giờ tăng ca (X+n, cột Giờ TC), TNC vẫn đúng; ngày công/giờ tăng ca hiện dạng số (không còn '24.').")

print("\nALL DONE")
