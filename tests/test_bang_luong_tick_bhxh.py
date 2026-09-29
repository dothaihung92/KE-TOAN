import os
import sys
import tempfile

_REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, _REPO_ROOT)
import server

# Yêu cầu người dùng (Bảng Lương): "chỗ BHXH thêm tick chọn lao động thì mới tính BHXH, còn không tích chọn
# thì không tính BHXH". Mỗi dòng nhân viên có ô tick "Đóng BHXH" (dong_bh): tick -> tính BH như cũ; bỏ tick ->
# BHXH/BHYT/BHTN (cả phần công ty lẫn phần người lao động) = 0, kéo theo thuế/thực lãnh tính lại.

TS = server._luong_chuan_tham_so(None)
BH_COT = ("bhxh_dn", "bhyt_dn", "bhtn_dn", "bhxh_nld", "bhyt_nld", "bhtn_nld")
dong = {"ma": "101", "ten": "Nguyễn A", "luong_cb": 10000000, "ngay_cong": 26, "tien_com": 730000,
        "muc_xang": 1000000, "muc_dt": 1000000, "trang_phuc": 400000, "so_npt": 0}

# ===== 1: có tick (hoặc dữ liệu cũ chưa có ô tick) -> tính BH y như trước; bỏ tick -> BH = 0. =====
co = server._luong_tinh_dong(dict(dong, dong_bh=1), TS)
cu = server._luong_tinh_dong(dict(dong), TS)                      # dữ liệu đã lưu trước khi có ô tick
khong = server._luong_tinh_dong(dict(dong, dong_bh=0), TS)
assert co["bhxh_dn"] == 1750000 and co["bhxh_nld"] == 800000 and co["bhyt_nld"] == 150000 and co["bhtn_nld"] == 100000
assert cu == co and cu["dong_bh"] == 1, "Dữ liệu cũ (chưa có ô tick) phải giữ nguyên kết quả"
assert all(khong[k] == 0 for k in BH_COT) and khong["bh_duoc_tru"] == 0, khong
print("PASS 1: tick -> tính BH; không tick -> BHXH/BHYT/BHTN (DN + NLĐ) đều 0; dữ liệu cũ giữ nguyên.")

# ===== 2: không đóng BH -> thu nhập tính thuế/thuế/thực lãnh tính lại đúng (không còn trừ BH NLĐ). =====
assert khong["tn_tinh_thue"] == co["tn_tinh_thue"] + co["bh_duoc_tru"]
assert khong["tt_luong"] == server._luong_lam_tron(khong["chi_phi_luong"] - khong["thue_tncn"])
assert abs(khong["kiem_tra"] - khong["thue_tncn"]) < 1e-6
assert khong["luong"] == co["luong"] and khong["chi_phi_luong"] == co["chi_phi_luong"]
print("PASS 2: bỏ tick -> thuế TNCN/thực lãnh tính lại đúng, lương/chi phí lương không đổi.")

# ===== 3: đọc giá trị tick từ nhiều dạng (checkbox, Excel, chữ). =====
c = server._luong_co_dong_bh
assert [c(v) for v in (1, True, "1", "x", "Có", "true", 2.0)] == [1] * 7
assert [c(v) for v in (0, False, "0", "false", "Không", "khong", 0.0)] == [0] * 7
assert c(None) == 1 and c("") == 1
print("PASS 3: nhận diện ô tick dạng 1/0, True/False, chữ Có/Không; trống = có (tương thích dữ liệu cũ).")

# ===== 4: /api/bang-luong-tinh + lưu/tải lại giữ nguyên trạng thái tick từng nhân viên. =====
import asyncio
import sqlite3


class _Req:
    def __init__(self, body): self._b = body
    async def json(self): return self._b


kq = asyncio.run(server.bang_luong_tinh(_Req({"rows": [dict(dong, dong_bh=1), dict(dong, ma="102", ten="B", dong_bh=0)]})))
r1, r2 = kq["rows"]
assert r1["bhxh_nld"] == 800000 and r2["bhxh_nld"] == 0 and r2["dong_bh"] == 0 and r1["dong_bh"] == 1

_duong_db = tempfile.mktemp(suffix=".sqlite3")


def _db_tam():
    c = sqlite3.connect(_duong_db)
    c.row_factory = sqlite3.Row
    return c


_goc_db = server.db
server.db = _db_tam
try:
    c0 = _db_tam()
    c0.execute("CREATE TABLE nhap_lieu (id INTEGER PRIMARY KEY AUTOINCREMENT, company_id INTEGER, loai TEXT, "
               "header_json TEXT, rows_json TEXT, updated_at TEXT, UNIQUE(company_id, loai))")
    c0.commit(); c0.close()
    asyncio.run(server.bang_luong_luu(1, _Req({"tham_so": {}, "thang": {"03": [dict(dong, dong_bh=1), dict(dong, ma="102", dong_bh=0)]}}), nam=2026))
    rows = server.bang_luong_get(1, nam=2026)["thang"]["03"]
    assert [r["dong_bh"] for r in rows] == [1, 0], rows
    assert rows[1]["bhxh_nld"] == 0 and rows[0]["bhxh_nld"] == 800000, rows
finally:
    server.db = _goc_db
print("PASS 4: lưu / tải lại giữ đúng tick từng nhân viên, BH của người không tick = 0.")

# ===== 5: xuất Excel ghi BH = 0 cho người không tick; nhập lại suy ra đúng tick (cả file mới xuất chưa có giá trị
# lưu sẵn lẫn file đã có giá trị). =====
import openpyxl
_goc_dl = server.DOWNLOAD_DIR
server.DOWNLOAD_DIR = tempfile.mkdtemp()
try:
    duong, _ = server._luong_xuat_excel(2026, TS, {"03": [dict(dong, dong_bh=1), dict(dong, ma="102", ten="B", dong_bh=0)]})
    ws = openpyxl.load_workbook(duong).active
    cot = [c[0] for c in server._LUONG_COT_EXCEL]
    ch = lambda r, k: ws.cell(r, cot.index(k) + 1).value
    assert str(ch(3, "bhxh_nld")).startswith("=") and ch(4, "bhxh_nld") == 0 and ch(4, "bhxh_dn") == 0
    t, loi = server._luong_doc_excel(openpyxl.load_workbook(duong, data_only=True), openpyxl.load_workbook(duong))
    assert loi == [] and [r["dong_bh"] for r in t["03"]] == [1, 0], t
finally:
    server.DOWNLOAD_DIR = _goc_dl
print("PASS 5: xuất/nhập Excel giữ đúng nhân viên nào đóng BH, nhân viên nào không.")

print("\nALL DONE")
