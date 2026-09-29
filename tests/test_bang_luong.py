import os
import sys
import json
import sqlite3
import tempfile

_REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, _REPO_ROOT)
import server

# Test cho tính năng BẢNG LƯƠNG (Nhập Liệu -> Bảng Lương, chọn theo năm) — xây theo file "TỔNG HỢP"
# (sheet TONG HOP VP) người dùng gửi: mỗi nhân viên 1 dòng / tháng, tính Tiền lương, BHXH/BHYT/BHTN (DN chịu +
# người lao động), Thu nhập chịu thuế, Thuế TNCN lũy tiến, TT Lương (thực lãnh), Chi phí lương.
# Các con số kiểm chứng bên dưới lấy từ GIÁ TRỊ EXCEL ĐÃ TÍNH SẴN trong chính file đó (không tự nghĩ ra).

TS = server._luong_chuan_tham_so(None)

# ===== 1: đọc số kiểu Việt Nam/gõ tay. =====
so = server._luong_so
assert so("32.500.000") == 32500000 and so("32,500,000") == 32500000 and so(32500000) == 32500000
assert so("26,5") == 26.5 and so("26.5") == 26.5 and so("1.234,5") == 1234.5 and so("1,234.5") == 1234.5
assert so("") == 0 and so(None) == 0 and so("abc") == 0 and so("-1.000") == -1000 and so(" 730.000 đ") == 730000
print("PASS 1: đọc số kiểu 32.500.000 / 32,500,000 / 26,5 / rỗng / chữ lạ.")

# ===== 2: làm tròn kiểu Excel (nửa ra xa số 0), không phải làm tròn chẵn của Python. =====
lt = server._luong_lam_tron
assert lt(0.5) == 1 and lt(1.5) == 2 and lt(2.5) == 3 and lt(123979.75) == 123980 and lt(-2.5) == -3 and lt(2.4) == 2
print("PASS 2: ROUND kiểu Excel (2,5 -> 3).")

# ===== 3: biểu thuế lũy tiến == công thức mảng gốc trong file: ROUND(SUM((x>={0,5,10,18,32,52,80}*1e6)*(x-{...}*1e6))*5%). =====
def thue_file(x):
    return sum((x >= t * 1e6) * (x - t * 1e6) for t in (0, 5, 10, 18, 32, 52, 80)) * 0.05
for x in (-5e6, 0, 1, 4999999, 5e6, 7e6, 10e6, 17999999, 19087500, 31e6, 51587500, 52e6, 79e6, 80e6, 120e6, 250e6):
    a = server._luong_thue_tncn(x, TS["bac_thue"])
    assert abs(a - thue_file(x)) < 1e-6, (x, a, thue_file(x))
print("PASS 3: biểu thuế lũy tiến khớp công thức mảng của file gốc (kể cả thu nhập âm -> 0).")

# ===== 4 (QUAN TRỌNG): tính 9 dòng THẬT trong file khớp Excel — TT Lương, Chi phí lương, thu nhập chịu thuế,
# không chịu thuế, BH được trừ, thu nhập tính thuế, thuế TNCN. =====
# (mã, lương CB, thưởng bán hàng, thưởng T13) + Excel: (TT Lương W, Chi phí X, chịu thuế AA, TN tính thuế AJ, thuế AK)
CA_THAT = [
    ("101", 32500000, 0, 0, 30050000, 35630000, 33500000, 19087500, 2167500),          # T7/8/9/10/11 NV 101 giống nhau
    ("102", 12000000, 1739595, 0, 15485615, 16869595, 14739595, 2479595, 123980),      # T10 NV 102
    ("102", 12000000, 1766284.1848677248, 0, 15510970, 16896284, 14766284.184867725, 2506284.184867725, 125314),  # T11
    ("101", 32500000, 0, 32500000, 55070625, 68130000, 66000000, 51587500, 9646875),   # T12 NV 101 (thưởng T13)
    ("102", 12000000, 1917779.9536317457, 12000000, 26339113, 29047780, 26917779.953631746, 14657779.953631744, 1448667),  # T12
]
for ma, luong, thuong, t13, w, x, aa, aj, ak in CA_THAT:
    k = server._luong_tinh_dong({"ma": ma, "luong_cb": luong, "ngay_cong": 26, "ngay_lam": 26, "tien_com": 730000,
                                 "muc_xang": 1000000, "muc_dt": 1000000, "trang_phuc": 400000,
                                 "thuong_bh": thuong, "thuong_t13": t13, "so_npt": 0}, TS)
    assert k["tt_luong"] == w and k["chi_phi_luong"] == x and abs(k["tn_chiu_thue"] - aa) < 0.01 \
        and abs(k["tn_tinh_thue"] - aj) < 0.01 and k["thue_tncn"] == ak, (ma, k)
    assert k["tn_khong_chiu_thue"] == 2130000 and k["giam_tru_ban_than"] == 11000000
    assert abs(k["kiem_tra"] - k["thue_tncn"]) < 1, ("Cột Kiểm tra số liệu phải ≈ Thuế TNCN", k["kiem_tra"], k["thue_tncn"])
print("PASS 4: 5 ca thật trong file (gồm thưởng bán hàng lẻ, thưởng T13, thuế 9.646.875đ) khớp Excel tới từng đồng.")

# ===== 5: các tình huống biên: người phụ thuộc, nghỉ bớt ngày, công chuẩn 0, thu nhập dưới giảm trừ, tham số tùy chỉnh. =====
co_ban = {"luong_cb": 26000000, "ngay_cong": 26, "ngay_lam": 26}
k0 = server._luong_tinh_dong(co_ban, TS)
k2 = server._luong_tinh_dong(dict(co_ban, so_npt=2), TS)
assert k2["tien_giam_tru_npt"] == 8800000 and k2["tn_tinh_thue"] == k0["tn_tinh_thue"] - 8800000
assert k2["thue_tncn"] < k0["thue_tncn"]
k_nghi = server._luong_tinh_dong(dict(co_ban, ngay_lam=13, muc_xang=1000000, muc_dt=600000), TS)
assert k_nghi["luong"] == 13000000 and k_nghi["xang_xe"] == 500000 and k_nghi["dien_thoai"] == 300000
assert k_nghi["bhxh_dn"] == 26000000 * 0.175, "BH tính theo lương CB, không theo ngày công thực tế (đúng file gốc)"
k_bang0 = server._luong_tinh_dong({"luong_cb": 1000000, "ngay_cong": 0, "ngay_lam": 0}, dict(TS, ngay_cong_chuan=0))
assert k_bang0["luong"] == 0
k_thap = server._luong_tinh_dong({"luong_cb": 5000000}, TS)
assert k_thap["thue_tncn"] == 0 and k_thap["tn_tinh_thue"] < 0
ts5 = server._luong_chuan_tham_so({"giam_tru_ban_than": "15.500.000", "giam_tru_npt": "6200000",
                                   "bac_thue": [[0, 5], [10000000, 10], [30000000, 20]]})
assert ts5["giam_tru_ban_than"] == 15500000 and ts5["bac_thue"] == [[0, 5], [10000000, 10], [30000000, 20]]
assert server._luong_thue_tncn(40000000, ts5["bac_thue"]) == 10e6 * .05 + 20e6 * .10 + 10e6 * .20
ts_hong = server._luong_chuan_tham_so({"bac_thue": "rác", "bh_dn": 5, "ngay_cong_chuan": "abc"})
assert ts_hong == TS, "Tham số rác phải rơi về mặc định, không crash"
print("PASS 5: người phụ thuộc, nghỉ bớt ngày (phụ cấp xăng/ĐT theo ngày), công chuẩn 0, thu nhập thấp, tham số riêng, tham số rác.")

# ===== 6: API lưu/đọc theo NĂM (DB tạm): tách riêng từng năm, kiểm tra năm hợp lệ, chuẩn hoá dữ liệu nhập. =====
import asyncio

class _Req:
    def __init__(self, body): self._b = body
    async def json(self): return self._b

_duong_db = tempfile.mktemp(suffix=".sqlite3")
def _db_tam():
    c = sqlite3.connect(_duong_db)
    c.row_factory = sqlite3.Row
    return c
_goc_db = server.db
server.db = _db_tam
try:
    c0 = _db_tam()
    c0.execute("CREATE TABLE nhap_lieu (id INTEGER PRIMARY KEY AUTOINCREMENT, company_id INTEGER, loai TEXT, "
               "header_json TEXT, rows_json TEXT, updated_at TEXT, UNIQUE(company_id, loai))")
    c0.commit(); c0.close()
    dong = {"ma": "101", "ten": "Phạm Ngọc Khánh", "luong_cb": "32.500.000", "ngay_cong": 26, "tien_com": "730000",
            "muc_xang": 1000000, "muc_dt": 1000000, "trang_phuc": 400000, "rac_khong_biet": "x"}
    kq = asyncio.run(server.bang_luong_luu(1, _Req({"tham_so": {}, "thang": {"07": [dong], "13": [dong]}}), nam=2025))
    assert kq["nam"] == 2025 and list(kq["thang"]) == ["07"], "Tháng 13 (ngoài 1-12) phải bị bỏ"
    r = kq["thang"]["07"][0]
    assert r["luong_cb"] == 32500000 and r["tt_luong"] == 30050000 and r["thue_tncn"] == 2167500 and "rac_khong_biet" not in r
    assert kq["cac_nam"] == [2025]
    # năm 2026 tách riêng, không lẫn dữ liệu 2025
    assert server.bang_luong_get(1, nam=2026)["thang"] == {}
    asyncio.run(server.bang_luong_luu(1, _Req({"tham_so": {"giam_tru_ban_than": 15500000},
                                               "thang": {"01": [dict(dong, luong_cb=20000000)]}}), nam=2026))
    g25, g26 = server.bang_luong_get(1, nam=2025), server.bang_luong_get(1, nam=2026)
    assert g25["tham_so"]["giam_tru_ban_than"] == 11000000 and g26["tham_so"]["giam_tru_ban_than"] == 15500000, (
        "Mỗi năm có tham số riêng")
    assert g25["thang"]["07"][0]["luong_cb"] == 32500000 and g26["thang"]["01"][0]["luong_cb"] == 20000000
    assert g26["cac_nam"] == [2026, 2025]
    # lưu lại năm = THAY THẾ (xoá dòng cũ)
    kq2 = asyncio.run(server.bang_luong_luu(1, _Req({"thang": {"08": [dong]}}), nam=2025))
    assert list(kq2["thang"]) == ["08"]
    for nam_sai in (1999, 2101, "abc"):
        try:
            server.bang_luong_get(1, nam=nam_sai)
            assert False, nam_sai
        except server.HTTPException as e:
            assert e.status_code == 400
    # công ty khác không thấy dữ liệu
    assert server.bang_luong_get(2, nam=2025)["thang"] == {}
    print("PASS 6: lưu/đọc theo năm (2025 ≠ 2026, tham số riêng từng năm), chuẩn hoá dữ liệu, chặn năm sai, tách theo công ty.")

    # ===== 7: dựng dòng lương từ 'Danh Sách Nhân Viên' (số kiểu 32.500.000, bỏ dòng không tên). =====
    hdr = ["STT", "Mã NV", "Họ và tên", "Chức vụ", "Lương Cơ bản", "PC Tiền cơm", "PC Xăng xe", "PC Chức vụ",
           "PC Điện thoại", "PC Trang phục"]
    rows = [[1, "101", "Phạm Ngọc Khánh", "Quản lý", "32.500.000", "730.000", "1.000.000", "", "1.000.000", "400.000"],
            [2, "", "", "Trống", "1", "1", "1", "1", "1", "1"],
            [3, "102", "Trần Minh Hùng", "Kinh doanh", 12000000, 730000, 1000000, 500000, 1000000, 400000]]
    c = _db_tam()
    c.execute("INSERT INTO nhap_lieu (company_id, loai, header_json, rows_json) VALUES (1,'nv',?,?)",
              (json.dumps(hdr), json.dumps(rows)))
    c.commit(); c.close()
    tv = server.bang_luong_tu_nhan_vien(1, nam=2025)["rows"]
    assert [x["ma"] for x in tv] == ["101", "102"], "Dòng không có tên phải bị bỏ"
    assert tv[0]["luong_cb"] == 32500000 and tv[0]["muc_xang"] == 1000000 and tv[0]["tien_com"] == 730000
    assert tv[1]["pc_chuc_vu"] == 500000 and tv[1]["ngay_cong"] == 26 and tv[1]["ghi_chu"] == "CK"
    assert server.bang_luong_tu_nhan_vien(2, nam=2025)["rows"] == []
    print("PASS 7: dựng dòng lương từ Danh Sách Nhân Viên (số kiểu VN, bỏ dòng trống, phụ cấp chức vụ được giữ).")

    # ===== 8: API tính lại trả cả tham số đã chuẩn hoá (giao diện dùng để hiển thị/lưu). =====
    kq8 = asyncio.run(server.bang_luong_tinh(_Req({"tham_so": {"giam_tru_ban_than": "11.000.000"},
                                                  "rows": [{"luong_cb": "32.500.000", "ngay_cong": 26}]})))
    assert kq8["tham_so"]["giam_tru_ban_than"] == 11000000 and kq8["rows"][0]["luong_cb"] == 32500000
    print("PASS 8: /api/bang-luong-tinh tính lại + trả tham số chuẩn hoá.")
finally:
    server.db = _goc_db
    try:
        os.remove(_duong_db)
    except Exception:
        pass

# ===== 9: xuất Excel CÓ CÔNG THỨC (như file gốc); nhập ngược file đó ra đúng dữ liệu; đọc file bố cục gốc. =====
import openpyxl
thang = {"07": [{"ma": "101", "ten": "Phạm Ngọc Khánh", "chuc_vu": "Quản lý", "luong_cb": 32500000, "ngay_cong": 26,
                 "muc_xang": 1000000, "muc_dt": 1000000, "tien_com": 730000, "trang_phuc": 400000, "ghi_chu": "CK"}],
         "12": [{"ma": "102", "ten": "Trần Minh Hùng", "chuc_vu": "Kinh doanh", "luong_cb": 12000000, "ngay_cong": 26,
                 "ngay_lam": 24, "muc_xang": 1000000, "muc_dt": 1000000, "tien_com": 730000, "trang_phuc": 400000,
                 "thuong_bh": 1917779.5, "thuong_t13": 12000000, "so_npt": 1, "ghi_chu": "CK"}]}
_goc_dl = server.DOWNLOAD_DIR
server.DOWNLOAD_DIR = tempfile.mkdtemp()
try:
    duong, ten_file = server._luong_xuat_excel(2025, TS, thang)
    assert ten_file == "BangLuong_2025.xlsx" and os.path.exists(duong)
    ws = openpyxl.load_workbook(duong).active
    cot = [c[0] for c in server._LUONG_COT_EXCEL]
    ch = lambda r, k: ws.cell(r, cot.index(k) + 1).value
    assert ch(3, "ten") == "Phạm Ngọc Khánh" and ch(3, "thang") == "07" and ch(4, "thang") == "12"
    assert str(ch(3, "luong")).startswith("=") and str(ch(3, "thue_tncn")).startswith("=ROUND(SUMPRODUCT(")
    assert "TỔNG CỘNG NĂM 2025" in str(ws.cell(5, 2).value)
    # dòng có ngày công thực tế 24 phải giữ nguyên trong file (không bị đổi về 26)
    assert ch(4, "ngay_lam") == 24 and ch(4, "so_npt") == 1
    # nhập ngược: file chưa có giá trị lưu sẵn (mới xuất từ phần mềm) vẫn đọc đủ, kể cả mức xăng xe/ĐT lấy từ công thức
    t2, loi2 = server._luong_doc_excel(openpyxl.load_workbook(duong, data_only=True), openpyxl.load_workbook(duong))
    assert loi2 == [], loi2
    assert set(t2) == {"07", "12"}
    a = t2["07"][0]
    assert (a["ma"], a["luong_cb"], a["muc_xang"], a["muc_dt"], a["tien_com"], a["ghi_chu"]) == \
           ("101", 32500000, 1000000, 1000000, 730000, "CK")
    b = t2["12"][0]
    assert (b["ngay_lam"], b["thuong_t13"], b["so_npt"], b["thuong_bh"]) == (24, 12000000, 1, 1917779.5)
    print("PASS 9: xuất Excel có công thức + dòng tổng; nhập ngược file xuất đọc lại đủ dữ liệu, bỏ dòng TỔNG CỘNG.")

    # file bố cục GỐC (giá trị đã tính sẵn, 2 dòng tiêu đề có ô gộp, cột A không tiêu đề, cột Tháng dạng chữ '07')
    wb = openpyxl.Workbook()
    w = wb.active
    w.title = "TONG HOP VP"
    w["B1"], w["C1"], w["D1"], w["E1"], w["G1"], w["H1"], w["I1"], w["W1"], w["Y1"], w["Z1"] = (
        "Họ và Tên", "Chức vụ", "Lương", "Ngày", "Tổng", "Tiền", "Phụ cấp", "TT Lương", "Tháng", "Ghi")
    w["D2"], w["E2"], w["G2"], w["H2"], w["I2"], w["J2"], w["L2"], w["M2"], w["N2"], w["O2"], w["P2"], w["Z2"], w["AH2"] = (
        "CB/Tháng", "Công", "NC", "Lương", "Tiền cơm", "Xăng xe", "Điện thoại", "Trang Phục", "THưởng Bán Hàng",
        "Thưởng T13", "ca", "chú", "Số người phụ thuộc")
    w["AH1"] = "Người"
    w["P1"] = "Tăng"
    w.merge_cells("I1:N1")
    w.merge_cells("W1:W2")
    w.merge_cells("Y1:Y2")
    w.append(["101", "Phạm Ngọc Khánh", "Quản lý", 32500000, 26, 0, 26, 32500000, 730000, 1000000, None, 1000000, 400000,
              None, None, None, None, None, None, None, None, None, 30050000, 35630000, "07", "CK"])
    w.append(["TỔNG CỘNG", None])
    w.append(["102", "Trần Minh Hùng", "Kinh doanh", 12000000, 26, 0, 26, 12000000, 730000, 1000000, None, 1000000,
              400000, 1739595, None, None, None, None, None, None, None, None, 15485615, 16869595, "07", "CK"])
    duong2 = os.path.join(server.DOWNLOAD_DIR, "goc.xlsx")
    wb.save(duong2)
    t3, loi3 = server._luong_doc_excel(openpyxl.load_workbook(duong2, data_only=True), None)
    assert loi3 == [] and len(t3["07"]) == 2, (t3, loi3)
    assert t3["07"][0]["muc_xang"] == 1000000 and t3["07"][0]["muc_dt"] == 1000000, "Mức phụ cấp suy ra từ giá trị đã tính"
    print("PASS 9b: đọc file theo bố cục gốc (2 dòng tiêu đề gộp ô, cột A không tiêu đề, tháng dạng chữ '07').")

    # lỗi dữ liệu: thiếu cột bắt buộc / tháng sai được báo rõ, không crash
    wb2 = openpyxl.Workbook()
    wb2.active.append(["Cột lạ", "Không phải bảng lương"])
    _t, _l = server._luong_doc_excel(wb2, None)
    assert _t == {} and _l, "File không phải bảng lương phải báo lỗi rõ"
    wb3 = openpyxl.load_workbook(duong2)
    wb3.active["Y3"] = "xx"
    wb3.active["Y5"] = 13
    wb3.save(duong2)
    t4, l4 = server._luong_doc_excel(openpyxl.load_workbook(duong2, data_only=True), None)
    assert t4 == {} and len(l4) == 2, (t4, l4)
    assert any("Tháng" in x for x in l4) and any("1-12" in x for x in l4), l4
    print("PASS 9c: file sai bố cục / tháng sai được báo lỗi rõ ràng, không crash.")

    # ===== 10: công thức trong file xuất tính lại độc lập == phần mềm (chỉ chạy nếu có thư viện 'formulas'). =====
    try:
        import warnings
        warnings.filterwarnings("ignore")
        import formulas
    except Exception:
        print("BỎ QUA Test 10: chưa cài thư viện 'formulas' (chỉ dùng để kiểm chứng công thức Excel độc lập)")
    else:
        sol = formulas.ExcelModel().loads(duong).finish().calculate()
        def gia_tri(r, khoa):
            from openpyxl.utils import get_column_letter as L
            ref = f"'[{os.path.basename(duong)}]TONG HOP'!{L(cot.index(khoa) + 1)}{r}".upper()
            for k, v in sol.items():
                if k.upper() == ref:
                    return float(v.value[0, 0])
            raise KeyError(ref)
        r = 3
        for t in ("07", "12"):
            k = server._luong_tinh_dong(thang[t][0], TS)
            for khoa in ("luong", "xang_xe", "dien_thoai", "tt_luong", "chi_phi_luong", "tn_chiu_thue",
                         "tn_khong_chiu_thue", "bh_duoc_tru", "tn_tinh_thue", "thue_tncn", "gio_tang_ca"):
                assert abs(gia_tri(r, khoa) - k[khoa]) < 0.01, (t, khoa, gia_tri(r, khoa), k[khoa])
            r += 1
        print("PASS 10: công thức Excel (gồm SUMPRODUCT biểu thuế) tính lại độc lập khớp phần mềm.")
finally:
    server.DOWNLOAD_DIR = _goc_dl

print("\nALL DONE")
