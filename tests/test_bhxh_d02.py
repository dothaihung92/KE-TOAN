import io
import os
import sys
import tempfile

_REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, _REPO_ROOT)
import openpyxl
import server

# Kê khai BHXH D02-LT: lập danh sách tăng/giảm/điều chỉnh từ Bảng Lương (so tháng này với tháng trước),
# xuất Excel (bảng chuẩn hoặc điền vào file mẫu tải từ cổng BHXH). Không tự điền thông tin thiếu.


def dong(ma, ten, luong, bh=1, cv="Nhân viên"):
    return {"ma": ma, "ten": ten, "luong_cb": luong, "dong_bh": bh, "chuc_vu": cv}


truoc = [dong("1", "Nguyễn Văn A", 5000000), dong("2", "Trần Thị B", 6000000), dong("3", "Lê C", 5500000),
         dong("4", "Phạm D", 7000000), dong("6", "Võ F", 5000000, bh=0)]
nay = [dong("1", "Nguyễn Văn A", 5000000),            # không đổi -> không có dòng
       dong("2-001", "Trần Thị B", 6500000),          # tăng lương (mã phiên bản lương)
       dong("4", "Phạm D", 6800000),                  # giảm lương
       dong("5", "Hồ E", 5200000),                    # tăng mới
       dong("6", "Võ F", 5000000),                    # bắt đầu đóng BHXH
       dong("7", "Đỗ G", 5000000, bh=0)]              # không đóng -> bỏ qua
# "3" Lê C không còn -> giảm hẳn
info = {("ma", "5"): {"cccd": "079000000005", "ngay_sinh": "01/02/1990", "gioi_tinh": "Nam"}}
kq = server.lap_d02_lt(truoc, nay, 2026, 5, info, da_tung_dong={("ma", "6")})
pa = {r["ten"]: r for r in kq}
assert set(pa) == {"Trần Thị B", "Phạm D", "Hồ E", "Võ F", "Lê C"}, set(pa)
assert pa["Hồ E"]["phuong_an"] == "TM" and pa["Hồ E"]["cccd"] == "079000000005" and pa["Hồ E"]["can_bo_sung"] == []
assert pa["Võ F"]["phuong_an"] == "ON"
assert pa["Trần Thị B"]["phuong_an"] == "TL" and pa["Trần Thị B"]["muc_luong"] == 6500000 and pa["Trần Thị B"]["muc_luong_cu"] == 6000000
assert pa["Trần Thị B"]["ma"] == "2"
assert pa["Phạm D"]["phuong_an"] == "GL"
assert pa["Lê C"]["phuong_an"] == "GH" and pa["Lê C"]["tu_thang"] == "05/2026"
assert set(pa["Lê C"]["can_bo_sung"]) == {"CCCD/ĐDCN", "Ngày sinh", "Giới tính"} and pa["Lê C"]["cccd"] == ""
assert [r["phuong_an"] for r in kq] == ["TM", "ON", "TL", "GL", "GH"]
print("PASS 1: tăng mới / đi làm lại / tăng-giảm lương / giảm hẳn đúng; thiếu thông tin thì để trống + báo.")

# 2) Xuất bảng chuẩn: bỏ dòng không chọn
kq[0]["chon"] = False
wb = openpyxl.load_workbook(io.BytesIO(server._bhxh_xuat_xlsx(kq, 2026, 5, "Cty X")))
ws = wb.active
assert [c.value for c in ws[3]] == server.BHXH_COT_D02
dl = [[c.value for c in r] for r in ws.iter_rows(min_row=4)]
assert len(dl) == 4 and dl[0][1] == "Võ F" and dl[0][11] == "ON" and dl[0][0] == 1
print("PASS 2: xuất Excel bảng chuẩn, bỏ dòng không chọn.")

# 3) Điền vào file mẫu tải từ cổng: khớp theo tên cột, bỏ qua dòng đánh số cột
mau = openpyxl.Workbook()
m = mau.active
m["A1"] = "MẪU D02-LT"
m.append([])
m.append(["STT", "Họ và tên", "Số CCCD", "Ngày sinh", "Mức lương", "Từ tháng", "Phương án", "Ghi chú"])
m.append(["(1)", "(2)", "(3)", "(4)", "(5)", "(6)", "(7)", "(8)"])
bio = io.BytesIO()
mau.save(bio)
kq[0]["chon"] = True
ws2 = openpyxl.load_workbook(io.BytesIO(server._bhxh_xuat_xlsx(kq, 2026, 5, "", bio.getvalue()))).active
assert ws2["A1"].value == "MẪU D02-LT" and ws2["A4"].value == "(1)"
assert [ws2.cell(5, c).value for c in range(1, 9)] == [1, "Hồ E", "079000000005", "01/02/1990", 5200000, "05/2026", "TM", "Tăng mới"]
assert ws2.cell(9, 2).value == "Lê C"
try:
    rong = openpyxl.Workbook()
    b2 = io.BytesIO()
    rong.save(b2)
    server._bhxh_xuat_xlsx(kq, 2026, 5, "", b2.getvalue())
    raise AssertionError("mẫu không có cột Họ và tên phải báo lỗi")
except ValueError:
    pass
print("PASS 3: điền vào file mẫu theo tên cột, giữ nguyên phần đầu mẫu.")

# 4) _bhxh_lap: đọc Bảng Lương 2 năm (tháng 1 so với tháng 12 năm trước), cảnh báo khi thiếu tháng trước
bl = {2025: {"12": truoc}, 2026: {"01": nay}}
goc = (server._luong_doc_nam, server.nhap_lieu_get)
server._luong_doc_nam = lambda cid, nam: ({}, bl.get(nam, {}), "", [])
server.nhap_lieu_get = lambda cid, loai="nv": {"header": ["Mã NV", "Họ và tên", "CCCD", "Ngày sinh"],
                                              "rows": [["3", "Lê C", "079000000003", "03/03/1985"]]}
try:
    rows, cb = server._bhxh_lap(1, 2026, 1)
    assert len(cb) == 1 and "thiếu thông tin" in cb[0], cb
    le = [r for r in rows if r["ten"] == "Lê C"][0]
    assert le["phuong_an"] == "GH" and le["cccd"] == "079000000003" and le["can_bo_sung"] == ["Giới tính"]
    rows2, cb2 = server._bhxh_lap(1, 2026, 2)
    assert rows2 == [] and "Chưa có Bảng Lương tháng 02/2026" in cb2[0]
    bl[2026]["03"] = nay
    rows3, cb3 = server._bhxh_lap(1, 2026, 3)
    assert all(r["phuong_an"] in ("TM", "ON") for r in rows3) and "tháng trước" in cb3[0], cb3
finally:
    server._luong_doc_nam, server.nhap_lieu_get = goc
print("PASS 4: lập theo Bảng Lương, nối sang năm trước; cảnh báo khi thiếu bảng lương.")

# 5) Lưu tài khoản (để trống mật khẩu = giữ mật khẩu cũ); đính kèm chỉ nhận file trong thư mục BHXH của công ty
tmp = tempfile.mkdtemp()
goc_db, goc_dd = server.DB_PATH, server.DATA_DIR
server.DB_PATH = os.path.join(tmp, "t.db")
server.DATA_DIR = tmp
try:
    server.init_db()
    conn = server.db()
    conn.execute("INSERT INTO companies (id, ten, mst) VALUES (9, 'Cty Chín', '0300000009')")
    conn.execute("INSERT INTO companies (id, ten, mst) VALUES (10, 'Cty Mười', '0300000010')")
    conn.commit()
    conn.close()
    server.bhxh_cau_hinh_luu(9, {"ten_dang_nhap": "TA1234", "ma_don_vi": "DV01", "mat_khau": "bi-mat"})
    server.bhxh_cau_hinh_luu(9, {"ten_dang_nhap": "TA1234", "mat_khau": ""})
    c = server.bhxh_cau_hinh_get(9)
    assert c["ten_dang_nhap"] == "TA1234" and c["ma_don_vi"] == "DV01" and c["co_mat_khau"] and "mat_khau" not in c
    # lưu chung với thông tin công ty: Sửa công ty thấy được, sửa công ty không gửi kèm / mật khẩu trống thì giữ nguyên
    ct = server.get_company_detail(9)
    assert ct["bhxh_user"] == "TA1234" and ct["bhxh_password"] == "bi-mat" and ct["bhxh_ma_don_vi"] == "DV01"
    server.update_company(9, {"ten": "Cty Chín", "mst": "0300000009"})
    ct = server.get_company_detail(9)
    assert ct["bhxh_user"] == "TA1234" and ct["bhxh_password"] == "bi-mat" and ct["bhxh_ma_don_vi"] == "DV01"
    server.update_company(9, {"ten": "Cty Chín", "mst": "0300000009", "bhxh_ma_don_vi": "DV02"})
    assert server.get_company_detail(9)["bhxh_ma_don_vi"] == "DV02"
    server.update_company(9, {"ten": "Cty Chín", "mst": "0300000009", "bhxh_user": "TA9999", "bhxh_password": ""})
    ct = server.get_company_detail(9)
    assert ct["bhxh_user"] == "TA9999" and ct["bhxh_password"] == "bi-mat"
    server.update_company(9, {"ten": "Cty Chín", "mst": "0300000009", "bhxh_user": "TA9999", "bhxh_password": "moi"})
    assert server._bhxh_tai_khoan(9)[:2] == ("TA9999", "moi")
    # dữ liệu bản trước (bảng bhxh_cfg) tự chuyển sang thông tin công ty
    conn = server.db()
    conn.execute("INSERT INTO bhxh_cfg (company_id, ma_don_vi, mat_khau) VALUES (10, 'CU10', 'pw10')")
    conn.commit()
    conn.close()
    assert server._bhxh_tai_khoan(10)[:2] == ("CU10", "pw10")
    assert server.get_company_detail(10)["bhxh_password"] == "pw10"
    r = server.bhxh_d02_xuat(9, {"nam": 2026, "thang": 5, "rows": [dict(x) for x in kq]})
    from urllib.parse import unquote
    p_luu = unquote(r.headers["x-duong-dan"])
    assert os.path.isfile(p_luu) and p_luu.endswith("D02-LT_05_2026.xlsx")
    assert openpyxl.load_workbook(io.BytesIO(r.body)).active["A2"].value == "Đơn vị: Cty Chín — Mã đơn vị: DV02"

    class _Drv:
        pass
    server.BHXH_DRIVERS[9] = _Drv()
    ngoai = os.path.join(tmp, "khac.xlsx")
    open(ngoai, "wb").write(b"x")
    try:
        server.bhxh_dinh_kem(9, {"duong_dan": ngoai})
        raise AssertionError("phải từ chối file ngoài thư mục BHXH")
    except server.HTTPException as e:
        assert e.status_code == 400
finally:
    server.BHXH_DRIVERS.pop(9, None)
    server.DB_PATH, server.DATA_DIR = goc_db, goc_dd
print("PASS 5: lưu tài khoản BHXH (không trả mật khẩu ra ngoài); đính kèm chỉ dùng file của công ty.")

html = open(os.path.join(_REPO_ROOT, "static", "index.html"), encoding="utf-8").read()
assert 'onclick="moBhxh()"' in html and "KHÔNG tự ký số, KHÔNG tự bấm Nộp" in html
assert 'id="m_bhxh_ma_dv"' in html and "bhxh_ma_don_vi:document.getElementById('m_bhxh_ma_dv')" in html
assert 'id="m_bhxh_user"' in html and 'id="m_bhxh_pass"' in html and "bhxh_password:document.getElementById('m_bhxh_pass')" in html
src = open(os.path.join(_REPO_ROOT, "server.py"), encoding="utf-8").read()
i0 = src.index("#  KÊ KHAI BHXH — D02-LT")
khoi = src[i0:src.index('if __name__ == "__main__":', i0)]
assert ".click()" not in khoi and "submit" not in khoi.lower(), "phần BHXH không được tự bấm nút/nộp"
print("PASS 6: giao diện có mục BHXH; mã không tự bấm Nộp.")

print("\nALL DONE")
