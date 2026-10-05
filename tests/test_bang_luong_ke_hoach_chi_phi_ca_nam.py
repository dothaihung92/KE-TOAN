import os
import sys
import asyncio
import sqlite3
import tempfile
import json
import random

_REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, _REPO_ROOT)
import server
from fastapi import HTTPException

# Yêu cầu người dùng (Bảng Lương): nút "Chi phí lương cả năm" — nhập khoảng tháng (vd T10–T12) + tổng chi phí lương
# (vd 50tr) -> phần mềm tự tính cần bao nhiêu người để tổng chi phí lương ĐÚNG BẰNG số đó; lương CB + phụ cấp giữ
# nguyên, thiếu thì tăng thưởng bán hàng + tăng ca; ưu tiên để người lao động dưới ngưỡng nộp thuế TNCN; người
# không đóng BHXH phải làm dưới 14 ngày/tháng và bị khấu trừ 10% thuế TNCN.

TS25 = server._luong_chuan_tham_so(None, 2025)
dong = {"ma": "1", "ten": "A", "luong_cb": 5_310_000, "tien_com": 700_000, "muc_xang": 500_000, "pc_chuc_vu": 500_000,
        "muc_dt": 500_000, "trang_phuc": 400_000}

# ===== 1: người không đóng BHXH + làm < 14 ngày bị khấu trừ 10% thuế TNCN (khi >= ngưỡng); còn lại tính như cũ. =====
tv = server._luong_tinh_dong(dict(dong, luong_cb=20_000_000, muc_xang=0, muc_dt=0, pc_chuc_vu=0, tien_com=0, trang_phuc=0,
                                  ngay_lam=13, dong_bh=0), TS25, "10")       # công chuẩn T10/2025 = 27 -> 13/27
assert tv["thoi_vu"] is True and tv["bhxh_nld"] == 0 and tv["giam_tru_ban_than"] == 0 and tv["tien_giam_tru_npt"] == 0
assert tv["tn_tinh_thue"] == tv["tn_chiu_thue"] and tv["thue_tncn"] == server._luong_lam_tron(tv["tn_chiu_thue"] * 0.1)
nho = server._luong_tinh_dong(dict(dong, ngay_lam=3, dong_bh=0), TS25, "10")        # thu nhập chịu thuế < 2 triệu -> không khấu trừ
assert nho["thoi_vu"] is True and nho["tn_chiu_thue"] < 2_000_000 and nho["thue_tncn"] == 0
lam14 = server._luong_tinh_dong(dict(dong, luong_cb=30_000_000, ngay_lam=14, dong_bh=0), TS25, "10")
assert lam14["thoi_vu"] is False and lam14["canh_bao_bh"] is True, "Làm >= 14 ngày phải đóng BHXH -> không phải thời vụ, cảnh báo"
dongbh13 = server._luong_tinh_dong(dict(dong, luong_cb=30_000_000, ngay_lam=13, dong_bh=1), TS25, "10")
assert dongbh13["thoi_vu"] is False and dongbh13["bhxh_nld"] > 0
assert server._luong_tinh_dong(dict(dong), TS25, "10")["canh_bao_bh"] is False
# ngưỡng khấu trừ theo năm
assert TS25["nguong_khau_tru_10"] == 2_000_000 and server._luong_chuan_tham_so(None, 2026)["nguong_khau_tru_10"] == 5_000_000
assert server._luong_chuan_tham_so({"nguong_khau_tru_10": "3.000.000"}, 2026)["nguong_khau_tru_10"] == 3_000_000
print("PASS 1: không đóng BHXH + <14 ngày -> khấu trừ 10% (từ ngưỡng); >=14 ngày không phải thời vụ mà bị cảnh báo.")

# ===== 2: kế hoạch: 7 nhân viên như Danh Sách Nhân Viên (5.310.000 + phụ cấp, chi phí đủ công 7.910.000/tháng). =====
pool = [server._luong_chuan_dong_nhap(dict(dong, ma=str(i), ten=f"NV{i}")) for i in range(2, 9)]


def kiem_tra_tong(th, muc):
    tong = sum(r["chi_phi_luong"] for rows in th.values() for r in rows)
    assert tong == muc, (tong, muc)
    return tong


# 2a: ví dụ người dùng: T10–T12, 50 triệu -> cần 2 người đủ công, tổng ĐÚNG 50 triệu, phần thiếu bù bằng thưởng + tăng ca, không thuế.
th, tom = server._luong_ke_hoach(pool, 2025, 10, 12, 50_000_000, None)
kiem_tra_tong(th, 50_000_000)
assert set(th) == {"10", "11", "12"} and tom["so_nguoi"] == 2 and tom["day_du"] == 2 and tom["thoi_vu"] == 0, tom
assert tom["tong_thue"] == 0 and tom["canh_bao"] == [] and tom["tong_thuong_bh"] > 0 and tom["tong_tang_ca"] > 0
for rows in th.values():
    for r in rows:
        assert r["dong_bh"] == 1 and r["ngay_lam_hd"] == r["ngay_cong_hd"], "Người đủ công + đóng BHXH"
        assert r["luong_cb"] == 5_310_000 and r["tien_com"] == 700_000, "Lương CB + phụ cấp giữ nguyên"
        assert r["thue_tncn"] == 0 and r["tn_tinh_thue"] <= 0, "Dưới ngưỡng nộp thuế TNCN"
        assert r["gio_tang_ca"] <= server._LUONG_GIO_TANG_CA_TOI_DA + 1e-9
print("PASS 2a: 50tr T10–T12 -> 2 người đủ công, tổng đúng 50.000.000, bù bằng thưởng + tăng ca, không thuế.")

# 2b: mục tiêu nhỏ hơn 1 người đủ công -> làm < 14 ngày, không BHXH, khấu trừ 10% (nếu >= ngưỡng).
th, tom = server._luong_ke_hoach(pool, 2025, 10, 12, 16_000_000, None)
kiem_tra_tong(th, 16_000_000)
assert tom["day_du"] == 0 and tom["thoi_vu"] >= 2 and tom["tong_thue"] == 0, tom     # nhiều người làm ít ngày để cùng dưới ngưỡng thuế
for rows in th.values():
    for r in rows:
        assert r["dong_bh"] == 0 and r["ngay_cong_hd"] - r["ngay_lam_hd"] >= 14 and r["bhxh_nld"] == 0 and r["bhxh_dn"] == 0, r["ngay_lam_hd"]
        assert r["thoi_vu"] is True
        if r["tn_chiu_thue"] >= 2_000_000:
            assert r["thue_tncn"] == server._luong_lam_tron(r["tn_chiu_thue"] * 0.1)
        else:
            assert r["thue_tncn"] == 0
print("PASS 2b: mục tiêu nhỏ -> nhiều người làm dưới 14 ngày, không BHXH, giữ dưới ngưỡng khấu trừ 10%; tổng đúng 16.000.000.")

# 2c: nhiều người đủ công + bù, phân bổ trong ngưỡng nộp thuế; tăng ca không vượt 40 giờ/tháng.
th, tom = server._luong_ke_hoach(pool, 2025, 10, 12, 200_000_000, None)
kiem_tra_tong(th, 200_000_000)
assert tom["so_nguoi"] == 7 and tom["day_du"] == 7 and tom["tong_thue"] == 0
th, tom = server._luong_ke_hoach(pool, 2025, 1, 12, 30_000_000, None, ty_le_tang_ca=100)   # 100% = CHỈ tăng ca (tối đa 40 giờ), KHÔNG dồn phần dư vào thưởng bán hàng
assert all(r["gio_tang_ca"] <= 40 + 1e-9 for rows in th.values() for r in rows) and tom["tong_thuong_bh"] == 0
tong100 = sum(r["chi_phi_luong"] for rows in th.values() for r in rows)
assert tong100 == 30_000_000 or any("THIẾU" in c for c in tom["canh_bao"]), "khớp mục tiêu hoặc báo còn thiếu (không dồn vào thưởng)"
th, tom = server._luong_ke_hoach(pool, 2025, 1, 12, 30_000_000, None, ty_le_tang_ca=60)    # dưới 100%: tăng ca + thưởng theo tỷ lệ, tổng khớp đúng
kiem_tra_tong(th, 30_000_000)
assert all(r["gio_tang_ca"] <= 40 + 1e-9 for rows in th.values() for r in rows)
print("PASS 2c: 200tr -> 7 người, không thuế; tăng ca không vượt 40 giờ/tháng; 100% không có thưởng bán hàng.")

# 2d: hết nhân viên mà vẫn thiếu -> dồn vào thưởng + cảnh báo (vượt ngưỡng thuế), tổng vẫn đúng.
th, tom = server._luong_ke_hoach(pool[:1], 2025, 10, 12, 300_000_000, None)
kiem_tra_tong(th, 300_000_000)
assert tom["canh_bao"] and tom["tong_thue"] > 0
print("PASS 2d: thiếu người trong danh sách -> có cảnh báo, tổng vẫn khớp.")

# 2e: khoảng tháng 1 tháng; các tháng khác trong năm đã có sẵn thì chỉ phân bổ phần còn lại.
th, tom = server._luong_ke_hoach(pool, 2025, 12, 12, 60_000_000, None, da_co_ngoai=44_000_000)
kiem_tra_tong(th, 16_000_000)
assert tom["can_them"] == 16_000_000 and tom["da_co_ngoai"] == 44_000_000 and set(th) == {"12"}
for sai in ((13, 12), (0, 5), (5, 4)):
    try:
        server._luong_ke_hoach(pool, 2025, sai[0], sai[1], 1_000_000, None)
        raise AssertionError("phải báo lỗi khoảng tháng")
    except HTTPException as e:
        assert e.status_code == 400
for muc, ngoai in ((0, 0), (10_000_000, 10_000_000), (10_000_000, 20_000_000)):
    try:
        server._luong_ke_hoach(pool, 2025, 1, 3, muc, None, da_co_ngoai=ngoai)
        raise AssertionError("phải báo lỗi mục tiêu")
    except HTTPException as e:
        assert e.status_code == 400
print("PASS 2e: trừ phần các tháng khác đã có; khoảng tháng/mục tiêu sai bị từ chối.")

# 2f: khớp đúng cả khi phụ cấp lẻ (làm tròn từng dòng) — thử nhiều mục tiêu ngẫu nhiên.
import random
random.seed(7)
for _ in range(40):
    muc = random.randrange(1_000_000, 120_000_000)
    tu = random.randrange(1, 12)
    th, tom = server._luong_ke_hoach(pool, 2026, tu, random.randrange(tu, 13), muc, None)
    kiem_tra_tong(th, muc)
    for rows in th.values():
        for r in rows:
            assert (r["dong_bh"] == 1) or (r["ngay_cong_hd"] - r["ngay_lam_hd"] >= 14), "Không đóng BHXH thì phải không lương từ 14 ngày làm việc"
            assert r["thuong_bh"] >= 0 and r["tang_ca"] >= 0
print("PASS 2f: 40 mục tiêu ngẫu nhiên đều khớp tổng đúng và không ai vừa làm >=14 ngày vừa không đóng BHXH.")

# ===== 3: API (DB tạm): đọc Danh Sách Nhân Viên của công ty, không lưu bảng lương. =====
class _Req:
    def __init__(self, body): self._b = body
    async def json(self): return self._b


_duong = tempfile.mktemp(suffix=".sqlite3")


def _db_tam():
    c = sqlite3.connect(_duong)
    c.row_factory = sqlite3.Row
    return c


_goc = server.db
server.db = _db_tam
try:
    c0 = _db_tam()
    c0.execute("CREATE TABLE nhap_lieu (id INTEGER PRIMARY KEY AUTOINCREMENT, company_id INTEGER, loai TEXT, "
               "header_json TEXT, rows_json TEXT, updated_at TEXT, UNIQUE(company_id, loai))")
    hdr = ["STT", "Mã NV", "Họ và tên", "Chức vụ", "Lương Cơ bản", "PC Tiền cơm", "PC Xăng xe", "PC Chức vụ", "PC Điện thoại", "PC Trang phục"]
    rows = [[i, str(i), f"NV{i}", "KD", 5310000, 700000, 500000, 500000, 500000, 400000] for i in range(2, 9)]
    c0.execute("INSERT INTO nhap_lieu (company_id, loai, header_json, rows_json) VALUES (1,'nv',?,?)", (json.dumps(hdr), json.dumps(rows)))
    c0.commit(); c0.close()
    kq = asyncio.run(server.bang_luong_ke_hoach(1, _Req({"nam": 2025, "tu_thang": 10, "den_thang": 12, "muc_tieu": "50.000.000",
                                                         "ty_le_tang_ca": 50, "tham_so": {}, "da_co_ngoai": 0})))
    assert kq["tom_tat"]["tong_chi_phi"] == 50_000_000 and kq["tom_tat"]["so_nguoi"] == 2 and set(kq["thang"]) == {"10", "11", "12"}
    assert sum(r["chi_phi_luong"] for rows_ in kq["thang"].values() for r in rows_) == 50_000_000
    try:
        asyncio.run(server.bang_luong_ke_hoach(2, _Req({"nam": 2025, "tu_thang": 1, "den_thang": 1, "muc_tieu": 1000000})))
        raise AssertionError("công ty chưa có nhân viên phải báo lỗi")
    except HTTPException as e:
        assert e.status_code == 400
    assert server.bang_luong_get(1, nam=2025)["thang"] == {}, "Xem trước không được tự lưu"
finally:
    server.db = _goc
print("PASS 3: API kế hoạch đọc Danh Sách Nhân Viên, trả bản xem trước đúng tổng, không tự lưu.")

# ===== 4: Excel: dòng thời vụ ghi thuế 10% + không giảm trừ; tính lại độc lập khớp phần mềm; nhập lại giữ nguyên. =====
import openpyxl
th, tom = server._luong_ke_hoach(pool, 2025, 10, 10, 16_000_000 // 3 * 1, None)
_dl = server.DOWNLOAD_DIR
server.DOWNLOAD_DIR = tempfile.mkdtemp()
try:
    dong_tv = dict(pool[0], ngay_lam=13, dong_bh=0, luong_cb=20_000_000, thuong_bh=1_000_000)
    dong_dh = dict(pool[1], dong_bh=1, so_npt=1)
    duong, _ = server._luong_xuat_excel(2025, TS25, {"10": [dong_tv, dong_dh]})
    ws = openpyxl.load_workbook(duong).active
    cot = [c[0] for c in server._LUONG_COT_EXCEL]
    ch = lambda r, k: ws.cell(r, cot.index(k) + 1).value
    assert "IF(" in ch(3, "thue_tncn") and "10%" in ch(3, "thue_tncn") and ch(3, "giam_tru_bt") == 0
    assert "SUMPRODUCT" in ch(4, "thue_tncn")
    t, loi = server._luong_doc_excel(openpyxl.load_workbook(duong, data_only=True), openpyxl.load_workbook(duong))
    assert loi == [] and [r["dong_bh"] for r in t["10"]] == [0, 1] and t["10"][0]["ngay_lam"] == 13
    try:
        import formulas
        sol = formulas.ExcelModel().loads(duong).finish().calculate()
        ten = lambda r, k: "'[%s]%s'!%s%d" % (os.path.basename(duong), ws.title.upper(), openpyxl.utils.get_column_letter(cot.index(k) + 1), r)
        mem = [server._luong_tinh_dong(dong_tv, TS25, "10"), server._luong_tinh_dong(dong_dh, TS25, "10")]
        for r_i, k in ((3, mem[0]), (4, mem[1])):
            assert round(float(list(sol[ten(r_i, "thue_tncn")].value[0])[0])) == k["thue_tncn"], (r_i, k["thue_tncn"])
            assert round(float(list(sol[ten(r_i, "tt_luong")].value[0])[0])) == k["tt_luong"], r_i
    except ImportError:
        pass
finally:
    server.DOWNLOAD_DIR = _dl
print("PASS 4: Excel: thời vụ có thuế 10% (công thức IF), người đóng BHXH giữ biểu lũy tiến; nhập lại đúng.")

# ===== 5: Người làm < 14 ngày (không BHXH): thuế 10% TRỪ VÀO thực lãnh (TT lương = Chi phí lương − thuế); tick "không trừ thuế
# 10%" thì thuế của họ = 0 và họ nhận đủ (TT lương = Chi phí lương). =====
tv2 = dict(dong, luong_cb=8_000_000, ngay_lam=13, dong_bh=0)
k_nld = server._luong_tinh_dong(tv2, TS25, "10")                     # mặc định: người lao động chịu (trừ vào thực lãnh)
assert TS25["thue_10_cong_ty_chiu"] is False
assert k_nld["thoi_vu"] and k_nld["thue_tncn"] > 0
assert k_nld["tt_luong"] == k_nld["chi_phi_luong"] - k_nld["thue_tncn"], (k_nld["tt_luong"], k_nld["chi_phi_luong"], k_nld["thue_tncn"])
assert k_nld["thue_tru_luong"] == k_nld["thue_tncn"] and abs(k_nld["kiem_tra"] - k_nld["thue_tncn"]) < 1
ts_ct = server._luong_chuan_tham_so({"thue_10_cong_ty_chiu": True}, 2025)
k_ct = server._luong_tinh_dong(tv2, ts_ct, "10")
assert k_ct["thue_tncn"] == 0 and k_ct["thue_tru_luong"] == 0 and k_ct["tt_luong"] == k_ct["chi_phi_luong"], "Không trừ thuế -> thuế về 0"
assert abs(k_ct["kiem_tra"]) < 1 and k_ct["thoi_vu"]
# người đóng BHXH không bị ảnh hưởng bởi tham số này
assert server._luong_tinh_dong(dict(dong, luong_cb=30_000_000, dong_bh=1), ts_ct, "10")["thue_tncn"] > 0
assert server._luong_chuan_tham_so({"thue_10_cong_ty_chiu": "không"}, 2025)["thue_10_cong_ty_chiu"] is False
assert server._luong_chuan_tham_so({"cong_ty_chiu_thue_10": True}, 2025)["thue_10_cong_ty_chiu"] is False, "Khóa cũ (mặc định bản trước) không còn hiệu lực"
# người đóng BHXH không bị ảnh hưởng: vẫn trừ BH NLĐ + thuế lũy tiến khỏi thực lãnh
dbh = server._luong_tinh_dong(dict(dong, luong_cb=30_000_000, dong_bh=1), TS25, "10")
assert dbh["tt_luong"] == dbh["chi_phi_luong"] - dbh["bh_duoc_tru"] - dbh["thue_tncn"] and dbh["thue_tru_luong"] == dbh["thue_tncn"]
# kế hoạch (mặc định): thực lãnh của người làm < 14 ngày = chi phí − thuế; bật "không trừ thuế" thì thuế = 0, thực lãnh = chi phí
pool_kh = [server._luong_chuan_dong_nhap(dict(dong, ma=str(i), ten=f"NV{i}", dong_bh=0)) for i in range(2, 9)]
th, tom = server._luong_ke_hoach(pool_kh, 2024, 10, 11, 60_000_000, None, rng=random.Random(1))
assert all(r["tt_luong"] == r["chi_phi_luong"] - r["thue_tncn"] for rs in th.values() for r in rs if r["thoi_vu"])
assert tom["tong_thue"] > 0
th, tom = server._luong_ke_hoach(pool_kh, 2024, 10, 11, 60_000_000, {"thue_10_cong_ty_chiu": True}, rng=random.Random(1))
assert all(r["tt_luong"] == r["chi_phi_luong"] and r["thue_tncn"] == 0 for rs in th.values() for r in rs if r["thoi_vu"]) and tom["tong_thue"] == 0
print("PASS 5: người làm < 14 ngày: thuế 10% trừ vào thực lãnh (mặc định); tick 'không trừ thuế' thì thuế = 0 và nhận đủ.")

# Excel khi tick "không trừ thuế 10%": dòng thời vụ có thuế =0 (công thức, nên nhập lại KHÔNG bị coi là chỉnh tay), TT lương = Chi phí lương
_dl = server.DOWNLOAD_DIR
server.DOWNLOAD_DIR = tempfile.mkdtemp()
try:
    duong, _ = server._luong_xuat_excel(2025, ts_ct, {"10": [tv2, dict(dong, luong_cb=30_000_000, dong_bh=1)]})
    ws = openpyxl.load_workbook(duong).active
    cot = [c[0] for c in server._LUONG_COT_EXCEL]
    ch = lambda r, k: ws.cell(r, cot.index(k) + 1).value
    assert ch(3, "thue_tncn") == "=0" and str(ch(4, "thue_tncn")).startswith("=ROUND(SUMPRODUCT")
    t, loi = server._luong_doc_excel(openpyxl.load_workbook(duong, data_only=True), openpyxl.load_workbook(duong))
    assert [r["thue_tay"] for r in t["10"]] == ["", ""] and loi == []
    try:
        import formulas
        sol = formulas.ExcelModel().loads(duong).finish().calculate()
        ten = lambda r, k: "'[%s]%s'!%s%d" % (os.path.basename(duong), ws.title.upper(), openpyxl.utils.get_column_letter(cot.index(k) + 1), r)
        a = server._luong_tinh_dong(tv2, ts_ct, "10")
        assert float(list(sol[ten(3, "thue_tncn")].value[0])[0]) == 0
        assert round(float(list(sol[ten(3, "tt_luong")].value[0])[0])) == a["tt_luong"] == a["chi_phi_luong"]
    except ImportError:
        pass
finally:
    server.DOWNLOAD_DIR = _dl

# ===== 6: Cột Thuế TNCN cho NGƯỜI DÙNG TỰ CHỈNH: thue_tay có số = dùng đúng số đó; trống = tự tính. =====
auto = server._luong_tinh_dong(tv2, TS25, "10")
tay = server._luong_tinh_dong(dict(tv2, thue_tay="300.000"), TS25, "10")
assert auto["thue_da_chinh"] is False and tay["thue_da_chinh"] is True and tay["thue_tncn"] == 300_000
assert tay["tt_luong"] == tay["chi_phi_luong"] - 300_000 and tay["thue_tru_luong"] == 300_000 and abs(tay["kiem_tra"] - 300_000) < 1
zero = server._luong_tinh_dong(dict(tv2, thue_tay=0), TS25, "10")
assert zero["thue_da_chinh"] is True and zero["thue_tncn"] == 0 and zero["tt_luong"] == zero["chi_phi_luong"], "Nhập 0 = miễn thuế"
assert server._luong_tinh_dong(dict(tv2, thue_tay=""), TS25, "10")["thue_tncn"] == auto["thue_tncn"], "Trống = tự tính"
dh = server._luong_tinh_dong(dict(dong, luong_cb=30_000_000, dong_bh=1, thue_tay=1_000_000), TS25, "10")      # người đóng BHXH cũng chỉnh được
assert dh["thue_tncn"] == 1_000_000 and dh["tt_luong"] == dh["chi_phi_luong"] - dh["bh_duoc_tru"] - 1_000_000
ct_chinh = server._luong_tinh_dong(dict(tv2, thue_tay=200_000), ts_ct, "10")                                  # đã tự chỉnh thì luôn trừ vào lương
assert ct_chinh["tt_luong"] == ct_chinh["chi_phi_luong"] - 200_000
# lưu / tải giữ số chỉnh tay
row_luu = server._luong_chuan_dong_nhap(dict(tv2, thue_tay="1.234.567"))
assert row_luu["thue_tay"] == 1_234_567 and server._luong_chuan_dong_nhap(dict(tv2))["thue_tay"] == ""
# Excel: số chỉnh tay ghi cố định, dòng tự tính giữ công thức; nhập lại nhận ra dòng nào đã chỉnh tay
_dl = server.DOWNLOAD_DIR
server.DOWNLOAD_DIR = tempfile.mkdtemp()
try:
    duong, _ = server._luong_xuat_excel(2025, TS25, {"10": [dict(tv2, thue_tay=300_000), dict(tv2, ma="9"), dict(dong, luong_cb=30_000_000, dong_bh=1)]})
    ws = openpyxl.load_workbook(duong).active
    cot = [c[0] for c in server._LUONG_COT_EXCEL]
    ch = lambda r, k: ws.cell(r, cot.index(k) + 1).value
    assert ch(3, "thue_tncn") == 300_000 and str(ch(4, "thue_tncn")).startswith("=IF(") and str(ch(5, "thue_tncn")).startswith("=ROUND(SUMPRODUCT")
    t, loi = server._luong_doc_excel(openpyxl.load_workbook(duong, data_only=True), openpyxl.load_workbook(duong))
    assert loi == [] and [r["thue_tay"] for r in t["10"]] == [300_000, "", ""], [r["thue_tay"] for r in t["10"]]
    try:
        import formulas
        sol = formulas.ExcelModel().loads(duong).finish().calculate()
        ten = lambda r, k: "'[%s]%s'!%s%d" % (os.path.basename(duong), ws.title.upper(), openpyxl.utils.get_column_letter(cot.index(k) + 1), r)
        assert round(float(list(sol[ten(3, "tt_luong")].value[0])[0])) == tay["tt_luong"]
        assert round(float(list(sol[ten(4, "tt_luong")].value[0])[0])) == server._luong_tinh_dong(dict(tv2, ma="9"), TS25, "10")["tt_luong"]
    except ImportError:
        pass
finally:
    server.DOWNLOAD_DIR = _dl
print("PASS 6: cột Thuế TNCN tự chỉnh được (trống = tự tính, 0 = miễn thuế); lưu/Excel giữ số đã chỉnh.")

print("\nALL DONE")
