import os, sys, io, base64, re
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import gpkd_soan as g
from PIL import Image

# (1) Quyết định của HĐTV (TNHH 2 thành viên đổi địa chỉ + đổi thành viên, CHƯA chọn Chủ tịch HĐTV): khối ký phải có HỌ TÊN (lấy người đại diện theo
#     pháp luật), không để trống. (2) Ảnh chữ ký / dấu + chữ ký phải VỪA khoảng trống phía trên họ tên và vừa bề rộng ô — không đè lên chữ.
def anh(w, h):
    b = io.BytesIO(); Image.new("RGBA", (w, h), (200, 0, 0, 255)).save(b, "PNG")
    return "data:image/png;base64," + base64.b64encode(b.getvalue()).decode()

def van_ban(d):
    t = [p.text for p in d.paragraphs]
    for tb in d.tables:
        for r in tb.rows:
            for c in r.cells:
                t += [p.text for p in c.paragraphs]
    return "\n".join(t)

du = {"loai": ["dia_chi", "chu_so_huu"], "dn": {"ten": "Công ty TNHH Biogreen", "ma_so": "0315253276", "loai_hinh": "Công ty trách nhiệm hữu hạn hai thành viên trở lên",
      "ngay_lap": "19/10/2026", "dia_danh": "TP. Hồ Chí Minh"},
      "hien_tai": {"dd_ho_ten": "Trần Khả Vy", "dia_chi": "1469/1A An Phú Đông", "ngay_dieu_le": "04/09/2018"},
      "dia_chi": {"so_nha": "Số 200 Đường Thạnh Xuân 52", "xa": "Phường Thới An", "tinh": "TP Hồ Chí Minh"},
      "thanh_vien": {"ht": [{"ho_ten": "Dương Ngọc Bảo Châu", "von": 2450000000, "cccd": "079197029102"}, {"ho_ten": "Dương Bảo Thịnh", "von": 2550000000, "bo": True, "cccd": "1"}],
                     "moi": [{"ho_ten": "Trần Khả Vy", "cccd": "093199002120", "von": 2550000000}]}}
d = g.soan_quyet_dinh(du, {"giam_doc": anh(400, 200)})
txt = van_ban(d)
i = txt.index("CHỦ TỊCH HỘI ĐỒNG THÀNH VIÊN")
assert "TRẦN KHẢ VY" in txt[i:], txt[i:i + 200]
assert "(Ký, ghi rõ họ và tên)" in txt[i:]
print("PASS 1: Quyết định HĐTV chưa chọn Chủ tịch -> khối ký ghi họ tên người đại diện (TRẦN KHẢ VY), không để trống.")

# người đại diện MỚI -> ghi tên người mới, KHÔNG gắn ảnh chữ ký của người đại diện hiện tại
du2 = dict(du, loai=["dia_chi", "chu_so_huu", "dai_dien"], dai_dien=[{"ho_ten": "Nguyễn Văn B"}])
d2 = g.soan_quyet_dinh(du2, {"giam_doc": anh(400, 200)})
t2 = van_ban(d2); i2 = t2.index("CHỦ TỊCH HỘI ĐỒNG THÀNH VIÊN")
assert "NGUYỄN VĂN B" in t2[i2:]
assert d2.element.xml.count("<wp:anchor") == 0, "không gắn chữ ký của người khác"
assert d.element.xml.count("<wp:anchor") == 1, "người đại diện hiện tại (đã có chữ ký trong Kho) -> gắn ảnh"
print("PASS 2: đổi người đại diện -> ghi tên người mới, không gắn chữ ký của người cũ.")

# (2) chữ ký TỰ CÂN CHỈNH: kích thước dấu/chữ ký không cố định -> ảnh co vừa bề rộng ô + khoảng chừa phía trên họ tên co/giãn theo đúng chiều cao ảnh
#     (ảnh nhỏ không chừa dòng trống thừa; ảnh lớn chừa vừa đủ; ảnh luôn nằm TRỌN trong khoảng chừa, không đè lên dòng "(Ký, ghi họ tên)" / họ tên)
def do(w, h, cot_cm):
    from docx import Document
    from docx.shared import Cm
    dd = Document(); tb = dd.add_table(rows=1, cols=1); c = tb.rows[0].cells[0]
    c.width = Cm(cot_cm)
    g._dong_trong_ky(c); pk = c.add_paragraph("TRẦN KHẢ VY")
    assert g._gan_chu_ky(pk, anh(w, h))
    x = pk._p.xml
    cx, cy = int(re.search(r'<wp:extent cx="(\d+)"', x)[1]), int(re.search(r'<wp:extent cx="\d+" cy="(\d+)"', x)[1])
    voff = int(re.search(r'<wp:positionV[^>]*>\s*<wp:posOffset>(-?\d+)<', x)[1])
    hoff = re.search(r'<wp:positionH[^>]*>\s*<wp:posOffset>(-?\d+)<', x)
    trong = 0
    e = pk._p.getprevious()
    while e is not None and e.tag.endswith("}p") and re.search(r'w:line="(\d+)"', e.xml):
        trong += int(re.search(r'w:line="(\d+)"', e.xml)[1]) / 20 * 12700
        e = e.getprevious()
    return cx, cy, voff, int(hoff[1]) if hoff else None, trong
import random
KHOANG_MAC_DINH = g._SO_DONG_KY * g._CAO_DONG_KY_PT * 12700
for w, h, cot in ((400, 200, 8.0), (800, 100, 8.0), (200, 400, 8.0), (400, 200, 4.0), (120, 60, 8.0), (3000, 1500, 8.0)):
    for _ in range(40):                                            # lặp để phủ độ lệch ngẫu nhiên
        cx, cy, voff, hoff, trong = do(w, h, cot)
        assert cx <= cot * 0.9 * 360000 + 500, ("rộng", w, h, cx)
        assert cy <= 4.0 * 360000 + 100, ("cao", w, h, cy)
        assert trong >= 1.2 * 360000 - 0.03 * 360000, ("khoảng chừa tối thiểu 1,2 cm", trong)
        assert voff <= -0.9 * trong + 0.1 * 360000 and voff >= -trong - 0.1 * 360000, ("ảnh bắt đầu từ đầu khoảng chừa (behind text)", w, h, voff, trong)
        assert trong <= max(cy * 0.8, 1.2 * 360000) + 0.03 * 360000, ("khoảng chừa không thừa: họ tên kéo sát lên", trong, cy)
        if hoff is not None:
            assert hoff >= 0 and hoff + cx <= cot * 360000 + 500, ("ảnh nằm trọn trong bề rộng ô", hoff, cx, cot)
_, cy_nho, _, _, tr_nho = do(800, 100, 8.0)        # chữ ký dẹt, thấp (vd chỉ có nét ký) -> ít chừa
_, cy_lon, _, _, tr_lon = do(200, 400, 8.0)        # dấu tròn/ảnh cao -> chừa nhiều hơn
assert tr_nho < tr_lon and tr_nho <= 1.2 * 360000 + 0.03 * 360000 and tr_lon > 2.5 * 360000, (tr_nho, tr_lon)
assert tr_lon < KHOANG_MAC_DINH * 1.6, "ảnh lớn cũng không chừa quá nhiều"
print("PASS 4: dấu/chữ ký tự co vừa bề rộng ô, khoảng chừa co/giãn theo chiều cao ảnh (ảnh nhỏ ít chừa, ảnh lớn đủ chỗ), ảnh luôn nằm trọn trong khoảng chừa.")

# (3) không có dòng trống thừa ở cuối văn bản / liên tiếp
from docx import Document
dd = Document(); dd.add_paragraph("Nội dung"); dd.add_paragraph(""); dd.add_paragraph(""); dd.add_paragraph("Tiếp"); dd.add_paragraph(""); dd.add_paragraph("")
dd.add_paragraph("Ký"); 
g._dong_trong_ky(dd)                                                   # khoảng chừa ký (cao cố định) KHÔNG bị coi là dòng trống thừa
dd.add_paragraph("HỌ TÊN"); dd.add_paragraph(""); dd.add_paragraph("")
g._bo_dong_trong_thua(dd)
vb = [p.text for p in dd.paragraphs]
assert vb[-1] == "HỌ TÊN" and vb.count("") == 2 + g._SO_DONG_KY, vb                 # mỗi chỗ có nhiều dòng trống liên tiếp chỉ còn 1 + 3 dòng chừa ký
print("PASS 5: bỏ dòng trống ở cuối + gộp dòng trống liên tiếp; giữ nguyên khoảng chừa chữ ký.")

# (4) bảng ký 2 bên (Giấy ủy quyền, Hợp đồng chuyển nhượng): khoảng chừa 2 ô BẰNG NHAU -> hai họ tên ngang hàng dù dấu bên này lớn hơn
def khoang_o(c):
    return sum(int(re.search(r'w:line="(\d+)"', e._p.xml)[1]) for e in c.paragraphs if not e.text.strip() and re.search(r'w:lineRule="exact"', e._p.xml))
for dd in (g.soan_giay_uy_quyen({"dn": {"ten": "CTY", "loai_hinh": "Công ty TNHH một thành viên"}, "hien_tai": {"dd_ho_ten": "A", "chu_so_huu": "A"}, "uy_quyen": {"ho_ten": "B"}},
                                {"ben_b": anh(900, 100), "giam_doc": anh(400, 400), "chu_so_huu": anh(400, 400)}),
           g.soan_hop_dong_chuyen_nhuong({"dn": {"ten": "CTY"}, "hien_tai": {}}, {"ho_ten": "Bán", "cccd": "079000000001"}, {"ho_ten": "Mua", "cccd": "079000000002"},
                                         {"cccd:079000000001": anh(900, 100), "cccd:079000000002": anh(400, 400)})):
    tb = [t for t in dd.tables if len(t.rows[0].cells) == 2][-1]
    a, b = (khoang_o(c) for c in tb.rows[0].cells)
    assert a == b and a > 0, (a, b)
    ten = [[p.text for p in c.paragraphs][-1] for c in tb.rows[0].cells]
    assert all(ten), ten
print("PASS 6: ký 2 bên: khoảng chừa hai ô bằng nhau (lấy bên lớn) -> họ tên hai bên ngang hàng; dấu to bên này không đẩy lệch.")
print("ALL DONE")
