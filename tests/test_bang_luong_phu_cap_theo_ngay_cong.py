import os
import sys
import tempfile

_REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, _REPO_ROOT)
import server

# Yêu cầu người dùng (Bảng Lương): "các phụ cấp phải tính theo ngày đi làm — đi làm đủ công mới nhận đủ phụ cấp,
# còn đi làm nghỉ vài ngày thì phụ cấp cũng sẽ giảm theo". Mọi phụ cấp = mức/tháng / công chuẩn x ngày đi làm.

TS = server._luong_chuan_tham_so(None)
dong = {"ma": "1", "ten": "A", "luong_cb": 10_000_000, "ngay_cong": 26, "tien_com": 730_000, "muc_xang": 1_000_000,
        "di_lai": 260_000, "pc_chuc_vu": 500_000, "muc_dt": 600_000, "trang_phuc": 400_000, "dong_bh": 0}

# ===== 1: đi làm ĐỦ công -> nhận đủ mức phụ cấp; nghỉ -> giảm theo tỷ lệ ngày đi làm. =====
du = server._luong_tinh_dong(dict(dong, ngay_lam=26), TS)
assert (du["tt_tien_com"], du["xang_xe"], du["tt_di_lai"], du["tt_pc_chuc_vu"], du["dien_thoai"], du["tt_trang_phuc"]) == \
       (730_000, 1_000_000, 260_000, 500_000, 600_000, 400_000)
nua = server._luong_tinh_dong(dict(dong, ngay_lam=13), TS)                # nghỉ nửa tháng
assert (nua["tt_tien_com"], nua["xang_xe"], nua["tt_di_lai"], nua["tt_pc_chuc_vu"], nua["dien_thoai"], nua["tt_trang_phuc"]) == \
       (365_000, 500_000, 130_000, 250_000, 300_000, 200_000)
nghi2 = server._luong_tinh_dong(dict(dong, ngay_lam=24), TS)              # nghỉ 2 ngày / 26
assert abs(nghi2["tt_tien_com"] - 730_000 / 26 * 24) < 1e-6 and abs(nghi2["tt_pc_chuc_vu"] - 500_000 / 26 * 24) < 1e-6
assert abs(nghi2["tt_trang_phuc"] - 400_000 / 26 * 24) < 1e-6
zero = server._luong_tinh_dong(dict(dong, ngay_lam=0), TS)                # nghỉ cả tháng -> không phụ cấp nào
assert all(zero[k] == 0 for k in ("tt_tien_com", "xang_xe", "tt_di_lai", "tt_pc_chuc_vu", "dien_thoai", "tt_trang_phuc"))
# ô Tổng NC trống -> đi làm đủ công chuẩn -> đủ phụ cấp
assert server._luong_tinh_dong(dict(dong), TS)["tt_tien_com"] == 730_000
print("PASS 1: đủ công nhận đủ phụ cấp; nghỉ 13/26 ngày còn một nửa; nghỉ 2 ngày giảm tỷ lệ; nghỉ cả tháng = 0.")

# ===== 2: phụ cấp đã giảm được đưa đúng vào thu nhập chịu thuế / không chịu thuế / tổng thu nhập. =====
assert nua["tn_khong_chiu_thue"] == 365_000 + 200_000 + 300_000                       # tiền cơm + trang phục + điện thoại
assert nua["tn_chiu_thue"] == 5_000_000 + 500_000 + 130_000 + 250_000                 # lương + xăng + đi lại + PC chức vụ
assert nua["chi_phi_luong"] == round(nua["tn_chiu_thue"] + nua["tn_khong_chiu_thue"])
assert du["tn_khong_chiu_thue"] == 730_000 + 400_000 + 600_000
print("PASS 2: phụ cấp đã giảm vào đúng thu nhập chịu thuế / không chịu thuế / chi phí lương.")

# ===== 3: công chuẩn theo lịch tháng: đủ công tháng 27 ngày thì nhận đủ, thiếu 1 ngày thì giảm 1/27. =====
ts25 = server._luong_chuan_tham_so(None, 2025)
d0 = dict(dong, ngay_cong=0)
t7 = server._luong_tinh_dong(dict(d0, ngay_lam=27), ts25, "07")     # công chuẩn T7/2025 = 27
assert t7["ngay_cong_hd"] == 27 and t7["tt_tien_com"] == 730_000
t7b = server._luong_tinh_dong(dict(d0, ngay_lam=26), ts25, "07")
assert abs(t7b["tt_tien_com"] - 730_000 * 26 / 27) < 1e-6
print("PASS 3: theo công chuẩn tháng (27 ngày): thiếu 1 ngày giảm 1/27 phụ cấp.")

# ===== 4: Excel xuất ra: mọi phụ cấp là công thức =(mức/công chuẩn)*ngày làm; nhập lại suy ra đúng mức. =====
import openpyxl
_dl = server.DOWNLOAD_DIR
server.DOWNLOAD_DIR = tempfile.mkdtemp()
try:
    duong, _ = server._luong_xuat_excel(2025, TS, {"03": [dict(dong, ngay_lam=13)]})
    ws = openpyxl.load_workbook(duong).active
    cot = [c[0] for c in server._LUONG_COT_EXCEL]
    ch = lambda k: ws.cell(3, cot.index(k) + 1).value
    for k in ("tien_com", "xang_xe", "di_lai", "pc_chuc_vu", "dien_thoai", "trang_phuc"):
        assert str(ch(k)).startswith("=(") and "*" in str(ch(k)), (k, ch(k))
    t, loi = server._luong_doc_excel(openpyxl.load_workbook(duong, data_only=True), openpyxl.load_workbook(duong))
    r = t["03"][0]
    assert loi == [] and (r["tien_com"], r["muc_xang"], r["di_lai"], r["pc_chuc_vu"], r["muc_dt"], r["trang_phuc"], r["ngay_lam"]) == \
           (730_000, 1_000_000, 260_000, 500_000, 600_000, 400_000, 13), r
    try:
        import formulas
        sol = formulas.ExcelModel().loads(duong).finish().calculate()
        ten = lambda k: "'[%s]%s'!%s3" % (os.path.basename(duong), ws.title.upper(), openpyxl.utils.get_column_letter(cot.index(k) + 1))
        for k, kv in (("tien_com", "tt_tien_com"), ("xang_xe", "xang_xe"), ("di_lai", "tt_di_lai"), ("pc_chuc_vu", "tt_pc_chuc_vu"),
                      ("dien_thoai", "dien_thoai"), ("trang_phuc", "tt_trang_phuc")):
            assert abs(float(list(sol[ten(k)].value[0])[0]) - nua[kv]) < 1e-6, (k, sol[ten(k)].value)
        assert round(float(list(sol[ten("tt_luong")].value[0])[0])) == nua["tt_luong"]
    except ImportError:
        pass
    # file cũ (theo mẫu gốc): tiền cơm/trang phục là số CỐ ĐỊNH, chưa tính theo ngày công -> đọc là mức
    wb = openpyxl.Workbook(); w = wb.active
    hdr1 = ["Mã NV", "Họ và Tên", "Chức vụ", "Lương", "Ngày", "Giờ", "Tổng", "Tiền", "Phụ cấp", "Phụ cấp", "Phụ cấp", "Tháng"]
    hdr2 = ["", "", "", "CB/Tháng", "Công", "Tăng ca", "NC", "Lương", "Tiền cơm", "Xăng xe", "Trang Phục", ""]
    w.append(hdr1); w.append(hdr2)
    w.append(["9", "B", "KD", 10_000_000, 26, 0, 13, 5_000_000, 730_000, 500_000, 400_000, 3])
    duong2 = os.path.join(server.DOWNLOAD_DIR, "cu.xlsx"); wb.save(duong2)
    t2, l2 = server._luong_doc_excel(openpyxl.load_workbook(duong2, data_only=True), None)
    r2 = t2["03"][0]
    assert r2["tien_com"] == 730_000 and r2["trang_phuc"] == 400_000 and r2["muc_xang"] == 1_000_000, (r2, l2)
finally:
    server.DOWNLOAD_DIR = _dl
print("PASS 4: Excel xuất ra dùng công thức theo ngày làm cho mọi phụ cấp, nhập lại đúng mức; file cũ số cố định đọc là mức.")

print("\nALL DONE")
