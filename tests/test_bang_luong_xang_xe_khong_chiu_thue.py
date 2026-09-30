import os
import sys
import tempfile

_REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, _REPO_ROOT)
import server
import openpyxl

# Yêu cầu: phụ cấp XĂNG XE là phụ cấp KHÔNG tính vào thu nhập chịu thuế TNCN (trước đây phần mềm cộng vào thu nhập chịu thuế).
TS = server._luong_chuan_tham_so(None, 2025)
d = {"luong_cb": 20_000_000, "ngay_cong": 26, "tien_com": 730_000, "di_lai": 260_000, "muc_dt": 500_000, "trang_phuc": 400_000, "dong_bh": 1}

# 1: tăng xăng xe -> thu nhập chịu thuế/tính thuế/thuế KHÔNG đổi; chỉ tăng thu nhập không chịu thuế, chi phí lương, thực lãnh
a = server._luong_tinh_dong(dict(d, muc_xang=0), TS, "05")
b = server._luong_tinh_dong(dict(d, muc_xang=3_000_000), TS, "05")
assert a["tn_chiu_thue"] == b["tn_chiu_thue"] == 20_000_000 + 260_000, "Chịu thuế = lương + đi lại (+ thưởng/tăng ca); không có xăng xe"
assert a["tn_tinh_thue"] == b["tn_tinh_thue"] and a["thue_tncn"] == b["thue_tncn"]
assert b["tn_khong_chiu_thue"] == a["tn_khong_chiu_thue"] + 3_000_000 == 730_000 + 3_000_000 + 400_000 + 500_000
assert b["chi_phi_luong"] == a["chi_phi_luong"] + 3_000_000 and b["tt_luong"] == a["tt_luong"] + 3_000_000
# xăng xe vẫn tính theo ngày đi làm
c = server._luong_tinh_dong(dict(d, muc_xang=2_600_000, ngay_lam=13), TS, "05")
assert c["xang_xe"] == 1_300_000 and c["tn_khong_chiu_thue"] == 365_000 + 1_300_000 + 200_000 + 250_000
print("PASS 1: xăng xe không chịu thuế: không đổi thu nhập chịu thuế/thuế; tăng thu nhập không chịu thuế + chi phí + thực lãnh; vẫn theo ngày làm.")

# 2: người làm thời vụ (khấu trừ 10%): xăng xe cũng không tính vào thu nhập chịu thuế để khấu trừ
tv = server._luong_tinh_dong(dict(d, dong_bh=0, ngay_lam=10, muc_xang=2_600_000), TS, "05")
assert tv["thoi_vu"] and abs(tv["tn_chiu_thue"] - (20_000_000 + 260_000) / 26 * 10) < 1
print("PASS 2: người làm < 14 ngày: xăng xe không nằm trong thu nhập bị khấu trừ 10%.")

# 3: Excel xuất: công thức chịu thuế không có cột xăng xe, không chịu thuế có; tính lại độc lập khớp phần mềm
_dl = server.DOWNLOAD_DIR
server.DOWNLOAD_DIR = tempfile.mkdtemp()
try:
    duong, _ = server._luong_xuat_excel(2025, TS, {"05": [dict(d, ma="1", ten="A", muc_xang=3_000_000, thuong_bh=1_000_000)]})
    ws = openpyxl.load_workbook(duong).active
    cot = [c[0] for c in server._LUONG_COT_EXCEL]
    L = lambda k: openpyxl.utils.get_column_letter(cot.index(k) + 1)
    chiu, khong = str(ws.cell(3, cot.index("tn_chiu_thue") + 1).value), str(ws.cell(3, cot.index("tn_khong_chiu_thue") + 1).value)
    assert L("xang_xe") + "3" not in chiu and L("xang_xe") + "3" in khong and L("tien_com") + "3" in khong and L("luong") + "3" in chiu
    try:
        import formulas
        sol = formulas.ExcelModel().loads(duong).finish().calculate()
        ten = lambda k: "'[%s]%s'!%s3" % (os.path.basename(duong), ws.title.upper(), L(k))
        k = server._luong_tinh_dong(dict(d, muc_xang=3_000_000, thuong_bh=1_000_000), TS, "05")
        for khoa in ("tn_chiu_thue", "tn_khong_chiu_thue", "tn_tinh_thue", "thue_tncn", "tt_luong", "chi_phi_luong"):
            kq = float(list(sol[ten(khoa)].value[0])[0])
            assert abs(kq - k[khoa]) < 1, (khoa, kq, k[khoa])
    except ImportError:
        pass
finally:
    server.DOWNLOAD_DIR = _dl
print("PASS 3: Excel xuất: xăng xe thuộc nhóm không chịu thuế; công thức tính lại khớp phần mềm.")

# 4: kế hoạch full công: xăng xe (khi có trần) được đẩy cùng nhóm phụ cấp không chịu thuế, không chiếm chỗ thu nhập chịu thuế
import random
base = {"luong_cb": 5_310_000, "tien_com": 700_000, "muc_xang": 500_000, "muc_dt": 500_000, "trang_phuc": 400_000}
pool = [server._luong_chuan_dong_nhap(dict(base, ma=str(i), ten=f"NV{i}", dong_bh=1)) for i in range(2, 9)]
th, tom = server._luong_ke_hoach(pool, 2025, 9, 12, 8_700_000 * 4, None, 50, 0, random.Random(1), True, {"muc_xang": "2000000"})
assert sum(r["chi_phi_luong"] for rows in th.values() for r in rows) == 8_700_000 * 4 and tom["tong_thuong_bh"] == 0 and tom["tong_tang_ca"] == 0, "Đủ chỗ trong phụ cấp không chịu thuế -> không cần thưởng/tăng ca"
assert all(r["muc_xang"] > 500_000 and r["tn_chiu_thue"] == r["luong"] for rows in th.values() for r in rows), "Xăng xe được đẩy; thu nhập chịu thuế chỉ còn lương"
print("PASS 4: kế hoạch full công đẩy xăng xe cùng nhóm phụ cấp không chịu thuế.")
print("\nALL DONE")
