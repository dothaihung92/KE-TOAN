import os
import sys
import asyncio
import sqlite3
import tempfile
import datetime

_REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, _REPO_ROOT)
import server

# Yêu cầu người dùng (Bảng Lương): "chỉnh công chuẩn 1 tháng cứ trừ chủ nhật ra ... tháng 30 ngày có 4 chủ
# nhật thì công là 26 ... trừ thêm các ngày nghỉ lễ trong tháng; ví dụ năm 2025 nghỉ lễ như thế nào thì tính
# công như vậy".

ct = server._luong_cong_chuan_thang

# ===== 1: ví dụ của người dùng: tháng 30 ngày có 4 Chủ nhật, không lễ -> 26 công. =====
# Tháng 6/2025: 30 ngày, CN = 1,8,15,22,29 (5 CN) -> 25 ; tháng 11/2025: 30 ngày, CN = 2,9,16,23,30 -> 25;
# tháng 4/2026: 30 ngày, CN = 5,12,19,26 (4 CN) -> 26 (chưa trừ lễ).
assert ct(2026, 4, []) == 26
assert ct(2025, 6, []) == 25 and ct(2025, 11, []) == 25
print("PASS 1: tháng 30 ngày có 4 Chủ nhật = 26 công; có 5 Chủ nhật = 25 công.")

# ===== 2: lịch lễ 2025 -> công chuẩn 12 tháng (đếm tay từ lịch). =====
le25 = server._luong_le_mac_dinh(2025)
ngay = {x["ngay"] for x in le25}
assert {"2025-01-01", "2025-01-27", "2025-01-31", "2025-04-07", "2025-04-30", "2025-05-01", "2025-09-01",
        "2025-09-02"} <= ngay and len(le25) == 11
for d in ngay:   # mọi ngày lễ mặc định phải đúng năm
    assert datetime.date.fromisoformat(d).year == 2025
cc = {t: ct(2025, t, le25) for t in server._LUONG_THANG}
# T1: 31 ngày - 4 CN (5,12,19,26) - lễ 1/1 (Tư) - 27..31/1 (5 ngày T2..T6) = 21
# T2: 28 - 4 CN = 24 ; T3: 31 - 5 CN (2,9,16,23,30) = 26 ; T4: 30 - 4 CN - 7/4 - 30/4 = 24
# T5: 31 - 4 CN - 1/5 = 26 ; T6: 30 - 5 CN = 25 ; T7: 31 - 4 CN = 27 ; T8: 31 - 5 CN (3,10,17,24,31) = 26
# T9: 30 - 4 CN (7,14,21,28) - 1/9 - 2/9 = 24 ; T10: 31 - 4 CN = 27 ; T11: 30 - 5 CN = 25 ; T12: 31 - 4 CN = 27
assert cc == {"01": 21, "02": 24, "03": 26, "04": 24, "05": 26, "06": 25, "07": 27, "08": 26, "09": 24,
              "10": 27, "11": 25, "12": 27}, cc
print("PASS 2: công chuẩn 12 tháng 2025 theo lịch (trừ Chủ nhật + lễ):", list(cc.values()))

# ===== 3: ngày lễ trùng Chủ nhật KHÔNG bị trừ 2 lần; ngày lễ ngoài tháng/năm bị bỏ. =====
assert ct(2026, 4, [{"ngay": "2026-04-26"}]) == 26            # 26/4/2026 là Chủ nhật
assert ct(2026, 4, [{"ngay": "2026-04-27"}]) == 25            # 27/4 (thứ Hai) trừ 1
assert server._luong_chuan_ngay_le([{"ngay": "2025-13-01"}, {"ngay": "2024-12-25"}, {"ngay": "2025-05-01", "ten": "x"},
                                    {"ngay": "2025-05-01"}, "rác"], 2025) == [{"ngay": "2025-05-01", "ten": "x"}]
print("PASS 3: lễ trùng Chủ nhật không trừ 2 lần; ngày lễ sai/ngoài năm/trùng bị lọc.")

# ===== 4: tham số có năm -> có cong_chuan; người dùng sửa lịch lễ thì công chuẩn đổi theo. =====
ts = server._luong_chuan_tham_so(None, 2025)
assert ts["cong_chuan"]["01"] == 21 and ts["ngay_le"][0]["ngay"] == "2025-01-01"
ts2 = server._luong_chuan_tham_so({"ngay_le": [{"ngay": "2025-01-02", "ten": "Nghỉ công ty"}]}, 2025)
assert ts2["cong_chuan"]["01"] == 26 and ts2["cong_chuan"]["09"] == 26        # lịch tự nhập THAY hẳn mặc định
assert server._luong_chuan_tham_so({"ngay_le": []}, 2025)["cong_chuan"]["01"] == 27
# năm chưa có lịch sẵn (2027): chỉ lễ dương lịch cố định, người dùng tự thêm Tết/Giỗ Tổ
assert [x["ngay"] for x in server._luong_le_mac_dinh(2027)][:2] == ["2027-01-01", "2027-04-30"]
print("PASS 4: tham số năm có công chuẩn từng tháng; sửa lịch lễ -> công chuẩn đổi; năm khác có lễ dương lịch.")

# ===== 5: tính lương theo công chuẩn của THÁNG: ô Ngày công trống/0 = theo lịch; nhập số = giữ số nhập. =====
dong = {"ma": "1", "ten": "A", "luong_cb": 10000000, "ngay_cong": 0}
t1 = server._luong_tinh_dong(dong, ts, "01")            # công chuẩn 21
t7 = server._luong_tinh_dong(dong, ts, "07")            # công chuẩn 27
assert t1["ngay_cong_hd"] == 21 and t7["ngay_cong_hd"] == 27 and t1["ngay_cong"] == 0
assert t1["luong"] == 10000000 and t7["luong"] == 10000000, "Đi làm đủ công chuẩn -> đủ lương cơ bản"
nghi = server._luong_tinh_dong(dict(dong, ngay_lam=20), ts, "01")     # nghỉ 1 ngày trong tháng 21 công
assert abs(nghi["luong"] - 10000000 / 21 * 20) < 1e-6 and nghi["ngay_lam_hd"] == 20
tay = server._luong_tinh_dong(dict(dong, ngay_cong=26), ts, "01")     # nhập tay 26 -> giữ 26
assert tay["ngay_cong_hd"] == 26 and tay["ngay_lam_hd"] == 26
cu = server._luong_tinh_dong(dong, server._luong_chuan_tham_so(None), "01")   # không có năm -> công chuẩn 26 như cũ
assert cu["ngay_cong_hd"] == 26
print("PASS 5: lương tính theo công chuẩn của từng tháng; nhập tay ngày công thì giữ số nhập.")

# ===== 6: API: dữ liệu đã lưu trước đây (ngày công bị điền cứng 26, chưa có lịch lễ) tự chuyển về 'theo lịch';
# tính lại/ lưu / tải giữ đúng; endpoint lịch lễ mặc định. =====
class _Req:
    def __init__(self, body): self._b = body
    async def json(self): return self._b


_duong = tempfile.mktemp(suffix=".sqlite3")


def _db_tam():
    c = sqlite3.connect(_duong)
    c.row_factory = sqlite3.Row
    return c


_goc = server.db
server.db = _db_tam
try:
    c0 = _db_tam()
    server._luong_dam_bao_bang(c0)
    import json
    cu_json = json.dumps({"01": [dict(dong, ngay_cong=26), dict(dong, ma="2", ngay_cong=24)]})
    c0.execute("INSERT INTO bang_luong (company_id, nam, tham_so_json, thang_json) VALUES (1, 2025, '{}', ?)", (cu_json,))
    c0.commit(); c0.close()
    g = server.bang_luong_get(1, nam=2025)
    a, b = g["thang"]["01"]
    assert a["ngay_cong"] == 0 and a["ngay_cong_hd"] == 21, a          # 26 cũ -> theo lịch
    assert b["ngay_cong"] == 24 and b["ngay_cong_hd"] == 24, b         # số nhập tay khác 26 giữ nguyên
    assert g["tham_so"]["nam"] == 2025 and g["tham_so"]["cong_chuan"]["01"] == 21
    # lưu lại cùng lịch lễ đã sửa -> tải lại đúng
    ts_luu = dict(g["tham_so"], ngay_le=[{"ngay": "2025-01-02", "ten": "Nghỉ"}])
    asyncio.run(server.bang_luong_luu(1, _Req({"tham_so": ts_luu, "thang": {"01": [dong]}}), nam=2025))
    g2 = server.bang_luong_get(1, nam=2025)
    assert g2["tham_so"]["ngay_le"] == [{"ngay": "2025-01-02", "ten": "Nghỉ"}]
    assert g2["thang"]["01"][0]["ngay_cong_hd"] == 26
    # tính lại từ giao diện: gửi năm + tháng
    kq = asyncio.run(server.bang_luong_tinh(_Req({"nam": 2025, "thang": "09", "tham_so": {}, "rows": [dong]})))
    assert kq["rows"][0]["ngay_cong_hd"] == 24 and kq["tham_so"]["cong_chuan"]["09"] == 24
    md = server.bang_luong_le_mac_dinh(2026)
    assert md["co_lich_day_du"] is True and md["cong_chuan"]["04"] == 24 and md["cong_chuan"]["02"] == 19, md["cong_chuan"]
    # nạp từ nhân viên: ngày công = 0 (theo lịch)
    r = server._luong_dong_tu_nhan_vien(["Họ và tên", "Lương Cơ bản"], [["A", 1000000]])
    assert r[0]["ngay_cong"] == 0
finally:
    server.db = _goc
print("PASS 6: dữ liệu cũ (26 cứng) chuyển về theo lịch; lưu/tải lịch lễ; tính lại theo năm+tháng; lịch lễ 2026.")

# ===== 7: xuất Excel dùng công chuẩn của từng tháng. =====
import openpyxl
_dl = server.DOWNLOAD_DIR
server.DOWNLOAD_DIR = tempfile.mkdtemp()
try:
    duong, _ = server._luong_xuat_excel(2025, ts, {"01": [dong], "07": [dong]})
    ws = openpyxl.load_workbook(duong).active
    cot = [c[0] for c in server._LUONG_COT_EXCEL]
    assert ws.cell(3, cot.index("ngay_cong") + 1).value == 21 and ws.cell(4, cot.index("ngay_cong") + 1).value == 27
finally:
    server.DOWNLOAD_DIR = _dl
print("PASS 7: Excel xuất ra ghi đúng công chuẩn từng tháng (T1=21, T7=27).")

print("\nALL DONE")
