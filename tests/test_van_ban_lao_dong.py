import os, sys, io, re, zipfile, datetime
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import van_ban_lao_dong as v
import xml.etree.ElementTree as ET

# Hợp đồng lao động / Quy chế lương / Thang bảng lương dựng từ Danh Sách Nhân Viên + Bảng Lương; xuất Word từ chính HTML người dùng đã sửa.
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

# --- số thành chữ (hợp đồng phải ghi lương bằng chữ)
assert v.doc_so_thanh_chu(5_250_000) == "Năm triệu hai trăm năm mươi nghìn đồng"
assert v.doc_so_thanh_chu(5_512_500) == "Năm triệu năm trăm mười hai nghìn năm trăm đồng"
assert v.doc_so_thanh_chu(15_000) == "Mười lăm nghìn đồng" and v.doc_so_thanh_chu(1_000_000) == "Một triệu đồng"
assert v.doc_so_thanh_chu(5_001) == "Năm nghìn không trăm lẻ một đồng" and v.doc_so_thanh_chu(21_000_000) == "Hai mươi mốt triệu đồng"
assert v.doc_so_thanh_chu(0) == "Không đồng" and v.doc_so_thanh_chu(1_250_000_000) == "Một tỷ hai trăm năm mươi triệu đồng"
# giới tính suy từ CCCD (chữ số thứ 4): 0/2 Nam, 1/3 Nữ; CMND 9 số -> không suy ra
assert v.gioi_tinh_tu_cccd("083190002386") == "Nữ" and v.gioi_tinh_tu_cccd("079090000123") == "Nam" and v.gioi_tinh_tu_cccd("025123456") == ""
# lương tối thiểu vùng theo năm
assert v.luong_toi_thieu_vung(2026, 1) == 5_310_000 and v.luong_toi_thieu_vung(2026, 2) == 4_730_000
assert v.luong_toi_thieu_vung(2025, 1) == 4_960_000 and v.luong_toi_thieu_vung(2023, 1) == 4_680_000
assert v.dia_danh_tu_dia_chi("174/21 Điện Biên Phủ, Phường 17, Quận Bình Thạnh, Thành phố Hồ Chí Minh, Việt Nam") == "TP. Hồ Chí Minh"
assert v.dia_danh_tu_dia_chi("12 Lê Lợi, Thành phố Huế") == "Huế"
assert v.dia_danh_tu_dia_chi("141/69/23,Khu Phố 3C") == "" and v.dia_danh_tu_dia_chi("") == "", "cụm cuối không phải thành phố/tỉnh -> không ghi địa danh"

HDR = ['STT', 'Mã NV', 'Họ và tên', 'Ngày sinh', 'Địa chỉ hiện đang cư trú', 'CCCD', 'Ngày cấp', 'Tháng/Năm vào làm', 'Đóng BHXH', 'Tháng/Năm nghỉ việc',
       'Chức vụ', 'Lương Cơ bản', 'PC Tiền cơm', 'PC Xăng xe', 'PC Điện thoại', 'PC Trang phục']
ROWS = [[1, 'NV1', 'Nguyễn Thị Chi', '06/09/1990', 'Bến Tre', '083190002386', '25/04/2021', '06/2025', 'x', '', 'Kế toán', 6000000, 730000, 500000, 300000, 0],
        [2, 'NV2', 'Trần Văn Bình', '01/01/1985', 'TP.HCM', '079085000111', '01/01/2021', '01/2024', 'x', '', 'Kế toán', 7_000_000, 730000, 500000, 300000, 0],
        [3, 'NV3', 'Lê Hoàng Cường', '', '', '', '', '', '', '', 'Giám đốc', 20_000_000, 730000, 0, 0, 0],
        [4, 'NV4', 'Phạm Dung', '02/02/1999', 'Q1', '079399000222', '', '02/2025', 'x', '12/2025', 'Tạp vụ', 4_000_000, 0, 0, 0, 0]]
BL = {"03": [{"ma": "NV1", "ten": "Nguyễn Thị Chi", "luong_cb": 6_500_000, "tien_com": 730000, "muc_xang": 500000, "muc_dt": 300000, "trang_phuc": 0, "di_lai": 200000, "ghi_chu": "CK"}],
      "07": [{"ma": "", "ten": "Người Mới", "chuc_vu": "Bảo vệ", "luong_cb": 6_000_000, "tien_com": 0, "muc_xang": 0, "muc_dt": 0, "trang_phuc": 0, "dong_bh": 1}]}
nv = v.gop_nhan_vien(HDR, ROWS, BL)
assert [n["ten"] for n in nv] == ["Nguyễn Thị Chi", "Trần Văn Bình", "Lê Hoàng Cường", "Phạm Dung", "Người Mới"] and [n["stt"] for n in nv] == [1, 2, 3, 4, 5]
a = nv[0]
assert a["luong_cb"] == 6_500_000 and a["di_lai"] == 200_000 and a["nguon"] == "Bảng lương tháng 3", "Bảng lương được ưu tiên"
assert any("6.000.000" in x and "6.500.000" in x for x in a["lech"]), a["lech"]
assert a["gioi_tinh"] == "Nữ" and nv[1]["gioi_tinh"] == "Nam" and nv[2]["gioi_tinh"] == ""
assert nv[1]["luong_cb"] == 7_000_000 and nv[1]["nguon"] == "Danh sách nhân viên"
assert nv[3]["da_nghi"] and not nv[0]["da_nghi"] and nv[2]["dong_bh"] is False and nv[0]["dong_bh"] is True
assert "chưa có trong Danh sách" in nv[4]["nguon"], "người chỉ có trong bảng lương vẫn được đưa vào"

CTY = {"ten": "CÔNG TY TNHH THỬ", "mst": "0300000001", "dia_chi": "1 Lê Lợi, Quận 1, Thành phố Hồ Chí Minh, Việt Nam", "nguoi_ky": "Hồ Thị Cẩm Vân"}
def van_ban(html):
    return re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", html).replace("&nbsp;", " ").replace("&amp;", "&"))

# --- hợp đồng lao động
tc = {"ong_ba_ky": "Bà", "loai_hd": "xdth", "so_thang": 12, "so_bat_dau": 7, "ngay_tra": 5, "dien_thoai": "028 3840 3379"}
h = v.dung_hop_dong(nv[0], CTY, tc, 0)
t = van_ban(h)
assert "Số: 07/HĐLĐ-2025" in t, "số HĐ theo mẫu + số bắt đầu"
assert "Bà: HỒ THỊ CẨM VÂN" in t and "Chức vụ: Giám đốc" in t and "Mã số thuế: 0300000001" in t
assert "Bà: NGUYỄN THỊ CHI" in t and "Giới tính: Nữ" in t and "083190002386" in t and "25/04/2021" in t and "Cục Cảnh sát quản lý hành chính" in t
assert "6.500.000 đồng/tháng" in t and "Sáu triệu năm trăm nghìn đồng" in t, "lương theo bảng lương + bằng chữ"
mucA = t.split("Phụ cấp lương và các khoản bổ sung khác:")[1].split("Hình thức trả lương")[0]
assert mucA.strip().startswith("Các khoản phụ cấp được tính theo số ngày công thực tế đi làm trong tháng") and "730.000" not in mucA and "đồng/tháng" not in mucA, "mặc định không ghi số tiền từng khoản phụ cấp"
tA = van_ban(v.dung_hop_dong(nv[0], CTY, dict(tc, ghi_phu_cap=True), 0))
assert "Phụ cấp tiền cơm (ăn trưa/ăn ca): 730.000" in tA and "Phụ cấp xăng xe, đi lại: 700.000" in tA and "điện thoại: 300.000" in tA and "Phụ cấp trang phục" not in tA, "tick ghi chi tiết: chỉ in phụ cấp > 0"
tK = van_ban(v.dung_hop_dong(nv[3], CTY, tc, 0))
assert "Phụ cấp lương và các khoản bổ sung khác: không có" in tK, "không có phụ cấp nào thì ghi không có"
assert "từ ngày 01 tháng 06 năm 2025 đến hết ngày 31 tháng 05 năm 2026" in t, "thời hạn 12 tháng tính từ ngày vào làm"
assert "ít nhất 30 ngày" in t, "HĐ xác định thời hạn 12-36 tháng: báo trước 30 ngày (Điều 35 BLLĐ)"
assert "mức tiền lương làm căn cứ đóng bảo hiểm là 6.500.000" in t
assert "TT số 21/2003" not in t and "45/2019/QH14" in t and "145/2020/NĐ-CP" in t, "bỏ căn cứ cũ, dùng BLLĐ 2019"
assert "07/HĐLĐ" in v.dung_hop_dong(nv[0], CTY, tc, 0) and "08/HĐLĐ-2024" in van_ban(v.dung_hop_dong(nv[1], CTY, tc, 1)), "số HĐ tăng dần"
h3 = van_ban(v.dung_hop_dong(nv[2], CTY, dict(tc, loai_hd="kxdth", bat_dau="01/03/2026", ngay_ky="28/02/2026"), 0))
assert "không xác định thời hạn" in h3 and "ít nhất 45 ngày" in h3 and "ngày 28 tháng 02 năm 2026" in h3 and "kể từ ngày 01 tháng 03 năm 2026" in h3
assert "Ông/Bà: LÊ HOÀNG CƯỜNG" in h3 and "đóng bảo hiểm là" not in h3, "không tick BHXH: không ghi mức đóng"
h6 = van_ban(v.dung_hop_dong(nv[0], CTY, dict(tc, so_thang=6, bat_dau="31/08/2025"), 0))
assert "đến hết ngày 28 tháng 02 năm 2026" in h6 and "ít nhất 03 ngày làm việc" in h6, "6 tháng: kết thúc 28/02, báo trước 3 ngày làm việc"
# năm lập hợp đồng chọn 2026: nhân viên vào làm từ năm trước -> ngày 01/01/2026 + số HĐ năm 2026; vào làm trong năm 2026 -> giữ ngày vào làm
assert v.ngay_bat_dau_theo_nam("12/2024", 2026) == datetime.date(2026, 1, 1) and v.ngay_bat_dau_theo_nam("", 2026) == datetime.date(2026, 1, 1)
assert v.ngay_bat_dau_theo_nam("15/03/2026", 2026) == datetime.date(2026, 3, 15) and v.ngay_bat_dau_theo_nam("06/2027", 2026) == datetime.date(2027, 6, 1)
h26 = van_ban(v.dung_hop_dong(nv[0], CTY, tc, 0, nam=2026))
assert "Số: 07/HĐLĐ-2026" in h26 and "ngày 01 tháng 01 năm 2026" in h26 and "2025" not in h26.split("Điều 1")[0], h26[:400]
assert "ngày 01 tháng 06 năm 2025" in van_ban(v.dung_hop_dong(nv[0], CTY, dict(tc, bat_dau="01/06/2025"), 0, nam=2026)), "người dùng nhập ngày bắt đầu thì giữ đúng"
assert "TP. Hồ Chí Minh, ngày" in van_ban(v.dung_hop_dong(nv[0], CTY, tc, 0, nam=2026))
h_kdd = van_ban(v.dung_hop_dong(nv[0], dict(CTY, dia_chi="141/69/23,Khu Phố 3C"), {"nguoi_ky": "A"}, 0, nam=2026))
assert "Khu Phố 3C, ngày" not in h_kdd and "ngày 01 tháng 01 năm 2026" in h_kdd, "không địa danh: chỉ hiện ngày tháng năm"
many = v.dung_hop_dong_nhieu(nv[:3], CTY, tc)
assert many.count('<section class="vb-trang">') == 3

# --- quy chế lương
qc = van_ban(v.dung_quy_che(nv[:4], CTY, {"ngay": "02/01/2026", "kem_phu_luc": True}, {"ngay_cong_chuan": 26}, 2026))
assert "293/2025/NĐ-CP" in qc and "41/2024/QH15" in qc and "110/2025/UBTVQH15" in qc and "5.310.000" in qc, "năm 2026: căn cứ văn bản mới"
# căn cứ pháp lý + lương tối thiểu theo NGÀY BAN HÀNH (không theo năm của Bảng Lương dùng làm số liệu)
assert "293/2025/NĐ-CP" in van_ban(v.dung_quy_che(nv[:4], CTY, {"ngay": "03/10/2026"}, {}, 2025)) and "5.310.000" in van_ban(v.dung_quy_che(nv[:4], CTY, {"ngay": "03/10/2026"}, {}, 2025))
qc25 = van_ban(v.dung_quy_che(nv[:4], CTY, {"ngay": "02/01/2025"}, {}, 2025))
assert "74/2024/NĐ-CP" in qc25 and "58/2014/QH13" in qc25 and "954/2020" in qc25 and "4.960.000" in qc25
assert "17,5%" in qc and "17.5%" not in qc and "8% bảo hiểm xã hội" in qc and "1,5% bảo hiểm y tế" in qc, "tỷ lệ BH lấy từ tham số bảng lương"
d5 = qc.split("Điều 5. Phụ cấp, trợ cấp, hỗ trợ")[1].split("Điều 6.")[0]
assert "tùy theo tính chất công việc" in d5 and "theo số ngày công thực tế đi làm trong tháng" in d5 and "pháp luật về thuế" in d5
assert "Kế toán" not in d5 and "730.000" not in d5 and "<table" not in v.dung_quy_che(nv[:4], CTY, {}, {}, 2026).split("Điều 6.")[0].split("Điều 5.")[1], "Điều 5 ghi chung theo quy định, không liệt kê bảng phụ cấp theo bảng lương"
assert "Giám đốc" in qc and "20.000.000" in qc and "PHỤ LỤC" in qc
assert "cao hơn bậc liền kề trước 5%." in van_ban(v.dung_thang_bang_luong(nv[:4], CTY, {}, 2026))
assert "cao hơn bậc liền kề trước 7,5%." in van_ban(v.dung_thang_bang_luong(nv[:4], CTY, {"buoc_pct": 7.5}, 2026)), "số thập phân kiểu Việt Nam (dấu phẩy)"
assert "Ví dụ" in qc and "÷ 26 × 25" in qc, "ví dụ tính lương lấy số thật"
m = re.search(r"Ví dụ: Chức danh (.+?) có lương cơ bản ([\d.]+) đồng, phụ cấp ([\d.]+) đồng.*?= ([\d.]+) đồng", qc)
assert m
cb_, pc_ = int(m.group(2).replace(".", "")), int(m.group(3).replace(".", ""))
assert int(m.group(4).replace(".", "")) == round((cb_ + pc_) / 26 * 25), "ví dụ tính đúng số học"
assert "150%" in qc and "200%" in qc and "300%" in qc and "40 giờ" in qc and "ngày 05 của tháng sau" in qc
assert "Thưởng cuối năm" in qc and "Hỗ trợ hiếu" not in qc, "khoản có số tiền mặc định không tick thì không đưa vào"
qc2 = van_ban(v.dung_quy_che(nv[:4], CTY, {"phuc_loi": {"hieu_hy": {"bat": True, "m1": 1000000, "m2": 500000}, "thuong_le": {"bat": True, "m1": 200000, "m2": 400000}}}, {}, 2026))
assert "bản thân người lao động 1.000.000 đồng" in qc2 and "từ 200.000 đồng đến 400.000 đồng" in qc2
assert "Phạm Dung" not in van_ban(v.dung_quy_che(nv[:4], CTY, {"kem_phu_luc": True}, {}, 2026)), "người đã nghỉ việc không đưa vào phụ lục"

# --- thang bảng lương
tl = v.tinh_thang_luong(nv[:4], 2026, vung=1, buoc_pct=5, so_bac=7)
g = {x["ten"]: x for x in tl["nhom"]}
assert tl["luong_toi_thieu"] == 5_310_000
kt = g["Kế toán"]["bac"]
assert kt[0] == 6_500_000 and kt[1] == round(6_500_000 * 1.05) and kt[6] == round(6_500_000 * 1.05 ** 6) and len(kt) == 7
xp = {x["nv"]["ten"]: x for x in g["Kế toán"]["xep"]}
assert xp["Nguyễn Thị Chi"]["bac"] == 1 and xp["Nguyễn Thị Chi"]["chenh"] == 0
assert xp["Trần Văn Bình"]["bac"] == 2 and xp["Trần Văn Bình"]["muc_bac"] == kt[1] and xp["Trần Văn Bình"]["chenh"] == 7_000_000 - kt[1]
assert g["Tạp vụ"]["bac"][0] == 5_310_000 and g["Tạp vụ"]["thap_hon_ltt"], "lương thấp hơn tối thiểu vùng: bậc 1 nâng lên bằng mức tối thiểu + cảnh báo"
assert tl["nhom"][0]["ten"] == "Giám đốc", "xếp nhóm lương cao trước"
cao = v.tinh_thang_luong([dict(nv[2], luong_cb=9_000_000), dict(nv[2], ten="B", luong_cb=14_000_000)], 2026, 1, 5, 7)["nhom"][0]
assert len(cao["bac"]) > 7 and cao["bac"][-1] >= 14_000_000 and not cao["vuot_bac_cuoi"], "có người vượt bậc 7 thì tự thêm bậc"
th = van_ban(v.dung_thang_bang_luong(nv[:4], CTY, {"ngay": "02/01/2026"}, 2026))
assert "HỆ THỐNG THANG LƯƠNG, BẢNG LƯƠNG" in th and "5.310.000" in th and "I II III IV V VI VII" in th.replace("  ", " ") and "Kế toán" in th and "Điều 93" in th
assert "BẢNG XẾP LƯƠNG HIỆN TẠI" in th and "Trần Văn Bình" in th

# --- đối chiếu
cb = v.kiem_tra(nv, 2026, {"vung": 1}, {"tien_com": 700000})
nd = " | ".join(c["nd"] for c in cb)
assert any(c["muc"] == "loi" and "Phạm Dung" in c["nd"] and "thấp hơn lương tối thiểu" in c["nd"] for c in cb)
assert "Lê Hoàng Cường: thiếu số CCCD" in nd and "Trần Văn Bình" in nd and "vượt trần 700.000" in nd
assert any("Lương cơ bản: Danh sách NV" in c["nd"] for c in cb) and any("Chức danh 'Kế toán'" in c["nd"] and "đi lại" in c["nd"] for c in cb)
assert v.kiem_tra([], 2026, {})[0]["muc"] == "loi"

# --- HTML -> Word
def doc_xml(data):
    z = zipfile.ZipFile(io.BytesIO(data))
    assert set(z.namelist()) >= {"[Content_Types].xml", "_rels/.rels", "word/document.xml", "word/styles.xml", "word/_rels/document.xml.rels"}
    for n in z.namelist():
        if not n.startswith("word/media/"):
            ET.fromstring(z.read(n))        # XML hợp lệ
    return ET.fromstring(z.read("word/document.xml")), z.read("word/styles.xml").decode("utf8")
def van_ban_docx(root):
    return ["".join(t.text or "" for t in p.iter(W + "t")) for p in root.iter(W + "p")]

data = v.html_sang_docx(many, {"font": "Arial", "size": 12, "line": 1.3, "le": [25, 25, 35, 20]})
root, st = doc_xml(data)
txt = "\n".join(van_ban_docx(root))
assert txt.count("HỢP ĐỒNG LAO ĐỘNG") == 3 and "NGUYỄN THỊ CHI" in txt and "TRẦN VĂN BÌNH" in txt and "Sáu triệu năm trăm nghìn đồng" in txt
assert len(list(root.iter(W + "pageBreakBefore"))) == 2, "3 hợp đồng = 2 lần sang trang"
assert 'w:ascii="Arial"' in st and 'w:val="24"' in st and 'w:line="312"' in st, "phông/cỡ/giãn dòng theo canh chỉnh"
pg = root.find(".//" + W + "pgMar")
assert pg.get(W + "left") == str(round(35 * 56.7)) and pg.get(W + "top") == str(round(25 * 56.7)) and pg.get(W + "right") == str(round(20 * 56.7))
assert len(list(root.iter(W + "tbl"))) == 6, "mỗi HĐ: bảng tiêu ngữ + bảng chữ ký"
assert any(j.get(W + "val") == "center" for j in root.iter(W + "jc")) and any(j.get(W + "val") == "both" for j in root.iter(W + "jc"))
assert len(list(root.iter(W + "b"))) > 20 and len(list(root.iter(W + "br"))) == 0
try:
    import docx as _docx
    d = _docx.Document(io.BytesIO(data))
    assert len(d.tables) == 6 and any("NGUYỄN THỊ CHI" in p.text for p in d.paragraphs)
except ImportError:
    pass

# người dùng SỬA TAY trên màn hình (contenteditable: <div>, <b>, style text-align...) -> Word giữ đúng nội dung sửa
sua = ('<section class="vb-trang"><div style="text-align: center;"><b>TIÊU ĐỀ MỚI</b></div><div>Đoạn  <i>nghiêng</i> và <u>gạch chân</u><br>dòng sau</div>'
       '<div><br></div><p class="j ti">Mục nội dung &amp; ký hiệu &lt;đặc biệt&gt;</p></section><section class="vb-trang"><p>Trang hai</p></section>')
root2, _ = doc_xml(v.html_sang_docx(sua))
ps = van_ban_docx(root2)
assert "TIÊU ĐỀ MỚI" in ps and "Mục nội dung & ký hiệu <đặc biệt>" in ps and "Trang hai" in ps and "Đoạn nghiêng và gạch chân" in "".join(ps).replace("  ", " ")
assert len(list(root2.iter(W + "i"))) == 1 and len(list(root2.iter(W + "u"))) == 1 and len(list(root2.iter(W + "br"))) == 1
assert len(list(root2.iter(W + "pageBreakBefore"))) == 1
assert [j.get(W + "val") for j in root2.iter(W + "jc")][:1] == ["center"]
ind = root2.find(".//" + W + "ind")
assert ind is not None and ind.get(W + "firstLine") == "567"

# bảng: tiêu đề cột (th) có đường viền + lặp đầu trang; colspan; khổ ngang
bang = ('<section class="vb-trang ngang"><table><colgroup><col style="width:30%"><col style="width:35%"><col style="width:35%"></colgroup>'
        '<tr><th>A</th><th colspan="2">B</th></tr><tr><td>x</td><td class="r">1.000</td><td class="c">y</td></tr></table></section>')
root3, _ = doc_xml(v.html_sang_docx(bang, {"ngang": True}))
tbl = root3.find(".//" + W + "tbl")
assert tbl.find(W + "tblPr/" + W + "tblBorders") is not None
cols = [int(c.get(W + "w")) for c in tbl.find(W + "tblGrid")]
assert len(cols) == 3 and abs(cols[0] / sum(cols) - 0.30) < 0.01 and abs(sum(cols) - (16838 - round(30 * 56.7) - round(15 * 56.7))) <= 3, cols
assert root3.find(".//" + W + "pgSz").get(W + "orient") == "landscape"
hang = list(tbl.iter(W + "tr"))
assert hang[0].find(W + "trPr/" + W + "tblHeader") is not None and hang[1].find(W + "trPr/" + W + "tblHeader") is None
assert tbl.find(".//" + W + "gridSpan").get(W + "val") == "2"
# bảng không viền (khối tiêu ngữ/chữ ký)
root4, _ = doc_xml(v.html_sang_docx('<table class="nb"><tr><td><p class="c">Trái</p></td><td><p class="c">Phải</p></td></tr></table>'))
assert root4.find(".//" + W + "tblBorders") is None
# HTML rỗng vẫn ra file hợp lệ
doc_xml(v.html_sang_docx(""))

# --- Excel thang bảng lương
import tempfile, openpyxl
p = os.path.join(tempfile.mkdtemp(), "tl.xlsx")
v.thang_luong_excel(p, nv[:4], CTY, {"ngay": "02/01/2026"}, 2026)
wb = openpyxl.load_workbook(p)
ws = wb["Thang bảng lương"]
vals = [[c for c in r if c is not None] for r in ws.iter_rows(values_only=True)]
assert any("HỆ THỐNG THANG LƯƠNG, BẢNG LƯƠNG" in str(r) for r in vals)
assert any(r and r[0] == "Mức lương" and r[1] == 6_500_000 and r[2] == round(6_500_000 * 1.05) for r in vals)
assert "Xếp lương" in wb.sheetnames and wb["Xếp lương"].max_row == 4
# --- thang bảng lương theo bố cục file mẫu
# nhóm tự khai báo: gộp nhiều chức danh, nhóm chưa có người dùng mức bậc 1 khai báo; người chưa thuộc nhóm nào tự lập nhóm
nhom_tc = "Giám đốc | 8000000\nKế toán; Tạp vụ\nPhó giám đốc; Kế toán trưởng"
tl2 = v.tinh_thang_luong(nv[:4], 2026, 1, 5, 7, nhom_tuy_chinh=nhom_tc)
ten2 = [g["ten"] for g in tl2["nhom"]]
assert ten2 == ["Giám đốc", "Kế toán; Tạp vụ", "Phó giám đốc; Kế toán trưởng"], ten2
assert tl2["nhom"][0]["bac"][0] == 8_000_000 and len(tl2["nhom"][0]["xep"]) == 1, "Giám đốc (Lê Hoàng Cường) vào nhóm khai báo, bậc 1 = 8.000.000"
assert len(tl2["nhom"][1]["xep"]) == 3 and tl2["nhom"][1]["bac"][0] == 5_310_000, "Kế toán + Tạp vụ gộp 1 nhóm; bậc 1 = mức thấp nhất (Tạp vụ 4tr nâng lên tối thiểu vùng)"
assert tl2["nhom"][2]["xep"] == [] and tl2["nhom"][2]["bac"][0] == 5_310_000, "nhóm trống: bậc 1 = lương tối thiểu vùng"
tl3 = v.tinh_thang_luong(nv[:3], 2026, 1, 5, 7, nhom_tuy_chinh="Giám đốc | 25000000")
assert tl3["nhom"][0]["bac"][0] == 25_000_000 and tl3["nhom"][0]["xep"][0]["bac"] == 0
assert tl3["nhom"][0]["duoi_bac_1"], "lương thấp hơn bậc 1 khai báo -> cờ cảnh báo"
cb3 = v.kiem_tra(nv[:3], 2026, {"vung": 1, "nhom_tuy_chinh": "Giám đốc | 25000000", "ngay": "02/01/2026"})
assert any("thấp hơn mức bậc 1" in c["nd"] for c in cb3), cb3
# hệ số lương = mức bậc / lương tối thiểu vùng
assert abs(tl["nhom"][0]["he_so"][0] - tl["nhom"][0]["bac"][0] / 5_310_000) < 1e-9
raw = v.dung_thang_bang_luong(nv[:4], CTY, {"ngay": "03/10/2026", "nhom_tuy_chinh": nhom_tc}, 2025)
th2 = van_ban(raw)
dau = th2.split("HỆ THỐNG THANG LƯƠNG")[0]
assert "Địa chỉ: 1 Lê Lợi" in dau and "ngày 03 tháng 10 năm 2026" not in dau, "đầu văn bản có địa chỉ; ngày nằm ở chỗ ký (như file mẫu)"
assert "Hệ số lương 1,51" in th2, "bậc 1 Giám đốc 8.000.000 / 5.310.000 = 1,51"
assert "1. Giám đốc" in th2 and "2. Kế toán; Tạp vụ" in th2 and "8.000.000" in th2 and "1,51" in th2
assert "5.310.000" in th2 and "293/2025" in th2 and "4.960.000" not in th2, "lương tối thiểu theo ngày ban hành (2026) dù lấy Bảng lương năm 2025"
assert "ngày 03 tháng 10 năm 2026 GIÁM ĐỐC CÔNG TY" in th2 and th2.index("(Ký, ghi rõ họ tên và đóng dấu)") < th2.index("PHỤ LỤC"), "ký ở cuối thang lương; xếp lương là phụ lục"
assert 'class="pb"' in raw and "Mức lương đóng BHXH" in th2 and "Không tham gia" in th2 and "Bảng lương năm 2025" in th2
assert "Hệ số lương" not in van_ban(v.dung_thang_bang_luong(nv[:4], CTY, {"hien_he_so": False}, 2026))
rootT, _ = doc_xml(v.html_sang_docx(raw, {"ngang": True, "size": 11}))
assert len(list(rootT.iter(W + "pageBreakBefore"))) == 1 and len(list(rootT.iter(W + "tbl"))) >= 4
p2 = os.path.join(tempfile.mkdtemp(), "tl2.xlsx")
v.thang_luong_excel(p2, nv[:4], CTY, {"ngay": "03/10/2026", "nhom_tuy_chinh": nhom_tc}, 2025)
w2 = openpyxl.load_workbook(p2)["Thang bảng lương"]
vv = [[c for c in r if c is not None] for r in w2.iter_rows(values_only=True)]
assert any(r and r[0] == "Hệ số lương" and r[1] == round(8_000_000 / 5_310_000, 2) for r in vv) and any(r and r[0] == "Mức lương" and r[1] == 8_000_000 for r in vv)
# --- chức danh chuẩn, dò lương theo năm, tiêu đề có năm
assert v.chuc_danh_tu_nhom(v.NHOM_MAC_DINH) == ["Giám đốc", "Phó giám đốc", "Kế toán trưởng", "Nhân viên kế toán", "Nhân viên kinh doanh", "Phân xưởng sản xuất"]
assert v.chuc_danh_tu_nhom("A; a | 100\nB") == ["A", "B"], "không trùng (không phân biệt hoa thường)"
assert v.bo_muc_bac_1("Giám đốc | 8000000\n\nPhó giám đốc; Kế toán trưởng") == "Giám đốc\nPhó giám đốc; Kế toán trưởng"
assert v.mac_dinh_tuy_chon(CTY)["nhom_tuy_chinh"] == v.NHOM_MAC_DINH
dl = v.do_luong_theo_nam(nv[:4], 2026, {"nhom_tuy_chinh": "Giám đốc | 99\nKế toán; Tạp vụ"})
ten_dl = [b["ten"] for b in dl["bang"]]
assert ten_dl == ["Giám đốc", "Kế toán; Tạp vụ"], ten_dl
assert dl["bang"][1]["so_nguoi"] == 2 and dl["bang"][1]["thap_nhat"] == 6_500_000 and dl["bang"][1]["cao_nhat"] == 7_000_000, "người đã nghỉ việc bị loại; dò thấp/cao nhất"
assert dl["nhom_tuy_chinh"] == "Giám đốc | 20000000\nKế toán; Tạp vụ | 6500000", dl["nhom_tuy_chinh"]
assert "HỆ THỐNG THANG LƯƠNG, BẢNG LƯƠNG NĂM 2025" in van_ban(v.dung_thang_bang_luong(nv[:4], CTY, {"ngay": "03/10/2026"}, 2025)), "tiêu đề ghi năm của Bảng lương"
# --- ảnh chữ ký trong HTML -> Word (hình nhúng, giữ tỷ lệ); gắn đúng ô chữ ký của người đó
import base64
from PIL import Image as _Im, ImageDraw as _Dr
_im = _Im.new("RGBA", (300, 100), (0, 0, 0, 0)); _Dr.Draw(_im).line([(10, 80), (80, 10), (150, 90), (290, 20)], fill=(0, 0, 128, 255), width=4)
_b = io.BytesIO(); _im.save(_b, "PNG"); URI = "data:image/png;base64," + base64.b64encode(_b.getvalue()).decode()
ck = {v.khoa_chu_ky("NV1", "Nguyễn Thị Chi"): URI, "giam_doc": URI}
hh = v.dung_hop_dong(nv[0], CTY, tc, 0, chu_ky=ck)
assert hh.count("<img") == 2 and "NGƯỜI SỬ DỤNG LAO ĐỘNG" in hh
assert v.dung_hop_dong(nv[1], CTY, tc, 0, chu_ky=ck).count("<img") == 1, "NV khác không có chữ ký riêng: chỉ chữ ký giám đốc"
assert v.dung_hop_dong(nv[0], CTY, dict(tc, gan_chu_ky=False), 0, chu_ky=ck).count("<img") == 0 and v.dung_hop_dong(nv[0], CTY, tc, 0).count("<img") == 0
rootI, _ = doc_xml(v.html_sang_docx(hh))
dr = list(rootI.iter(W + "drawing"))
assert len(dr) == 2
ext = dr[0].find(".//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent")
assert int(ext.get("cy")) == 80 * 9525 and abs(int(ext.get("cx")) / int(ext.get("cy")) - 3.0) < 0.02, "cao 80px, đúng tỷ lệ ảnh"
zi = zipfile.ZipFile(io.BytesIO(v.html_sang_docx(hh)))
assert [n for n in zi.namelist() if n.startswith("word/media/")] == ["word/media/chuky1.png", "word/media/chuky2.png"] and "image/png" in zi.read("[Content_Types].xml").decode()
assert 'Target="media/chuky2.png"' in zi.read("word/_rels/document.xml.rels").decode()
# src không phải ảnh data URI (vd đường dẫn ngoài) bị bỏ qua, không làm hỏng file
doc_xml(v.html_sang_docx('<p>a <img src="http://x/y.png"> b</p>'))
print("PASS")
