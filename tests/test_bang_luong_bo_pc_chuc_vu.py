import os
import sys
import tempfile

_REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, _REPO_ROOT)
import server
import openpyxl

# Yêu cầu: "bỏ cột phụ cấp chức vụ luôn" khỏi Bảng Lương (phụ cấp chức vụ có tính BHXH nên không đưa vào bảng lương này).

TS = server._luong_chuan_tham_so(None, 2025)

# 1: dòng có pc_chuc_vu (dữ liệu cũ/API) -> bị bỏ qua, không cộng vào lương/chi phí/thu nhập chịu thuế
a = server._luong_tinh_dong({"luong_cb": 6_000_000, "ngay_cong": 26, "dong_bh": 1}, TS, "05")
b = server._luong_tinh_dong({"luong_cb": 6_000_000, "ngay_cong": 26, "dong_bh": 1, "pc_chuc_vu": 500_000}, TS, "05")
assert a["chi_phi_luong"] == b["chi_phi_luong"] == 6_000_000 and a["tn_chiu_thue"] == b["tn_chiu_thue"] and a["tt_luong"] == b["tt_luong"]
assert "pc_chuc_vu" not in b and "tt_pc_chuc_vu" not in b and "pc_chuc_vu" not in server._luong_chuan_dong_nhap({"pc_chuc_vu": 5})
print("PASS 1: phụ cấp chức vụ không còn tính vào bảng lương (kể cả dữ liệu cũ có sẵn số này).")

# 2: nạp từ Danh Sách Nhân Viên: bỏ qua cột PC Chức vụ
hd = ["Mã NV", "Họ và tên", "Chức vụ", "Lương Cơ bản", "PC Tiền cơm", "PC Xăng xe", "PC Chức vụ", "PC Điện thoại", "PC Trang phục"]
rows = server._luong_dong_tu_nhan_vien(hd, [["1", "A", "KD", 5_310_000, 700_000, 500_000, 500_000, 500_000, 400_000]])
assert rows[0]["muc_xang"] == 500_000 and rows[0]["muc_dt"] == 500_000 and "pc_chuc_vu" not in rows[0]
assert server._luong_tinh_dong(rows[0], TS, "05")["chi_phi_luong"] == 5_310_000 + 700_000 + 500_000 + 500_000 + 400_000
print("PASS 2: nạp từ Danh Sách Nhân Viên không lấy PC Chức vụ.")

# 3: Excel xuất không còn cột PC Chức vụ; đọc lại file cũ có cột này không lỗi; tổng khớp
_dl = server.DOWNLOAD_DIR
server.DOWNLOAD_DIR = tempfile.mkdtemp()
try:
    assert not any(c[0] == "pc_chuc_vu" for c in server._LUONG_COT_EXCEL)
    duong, _ = server._luong_xuat_excel(2025, TS, {"05": [dict(rows[0], ngay_lam="")]})
    ws = openpyxl.load_workbook(duong).active
    tieu = " ".join(str(c.value or "") for c in ws[1]) + " " + " ".join(str(c.value or "") for c in ws[2])
    assert "PC Chức vụ" not in tieu
    t, loi = server._luong_doc_excel(openpyxl.load_workbook(duong, data_only=True), openpyxl.load_workbook(duong))
    assert loi == [] and t["05"][0]["muc_xang"] == 500_000 and "pc_chuc_vu" not in t["05"][0]
finally:
    server.DOWNLOAD_DIR = _dl
print("PASS 3: Excel không còn cột PC Chức vụ, đọc lại bình thường.")
# 4: Danh Sách Nhân Viên: bộ cột chuẩn không còn PC Chức vụ; import file có cột "Phụ cấp chức vụ" thì bỏ cột đó và KHÔNG làm sai cột Chức vụ
import asyncio, io
assert "PC Chức vụ" not in server.NV_HEADERS and len(server.NV_HEADERS) == 17 and "Tháng/Năm nghỉ việc" in server.NV_HEADERS and "Tháng/Năm thay đổi lương" in server.NV_HEADERS
wb = openpyxl.Workbook()
w = wb.active
for pos, dong in enumerate([["Mã NV", "Họ và tên", "Phụ cấp chức vụ", "Chức vụ", "Lương cơ bản", "Tiền cơm", "Xăng xe", "Điện thoại", "Trang phục"],
                            [],
                            ["1", "Nguyễn A", 900_000, "Kinh doanh", 5_310_000, 700_000, 500_000, 500_000, 400_000]]):
    w.append(dong)
buf = io.BytesIO()
wb.save(buf)

class Up:
    filename = "nv.xlsx"
    async def read(self): return buf.getvalue()

class Form:
    def getlist(self, k): return [Up()] if k == "files" else []
    def get(self, k): return None

class Req:
    async def form(self): return Form()

kq = asyncio.run(server.nhap_lieu_import_nhan_vien(1, Req()))
assert kq["header"] == server.NV_HEADERS and len(kq["rows"]) == 1
r = dict(zip(kq["header"], kq["rows"][0]))
assert r["Chức vụ"] == "Kinh doanh" and r["Lương Cơ bản"] == 5_310_000 and r["PC Tiền cơm"] == 700_000 and r["PC Xăng xe"] == 500_000
assert 900_000 not in kq["rows"][0], "Phụ cấp chức vụ của file nguồn không được đưa vào bất kỳ cột nào"
print("PASS 4: import Danh Sách Nhân Viên bỏ cột phụ cấp chức vụ, không nhầm sang cột Chức vụ.")

# 5: cột "Tháng/Năm nghỉ việc": nghỉ tháng M -> còn lên bảng lương tới hết tháng M, từ tháng M+1 không còn; trống = còn làm
hd = ["Mã NV", "Họ và tên", "Tháng/Năm nghỉ việc", "Lương Cơ bản"]
nv = [["1", "Còn làm", "", 5_000_000], ["2", "Nghỉ 6/2025", "06/2025", 5_000_000], ["3", "Nghỉ ngày đầy đủ", "15/03/2025", 5_000_000], ["4", "Nghỉ 12/2024", "2024-12", 5_000_000]]
ten = lambda nam, thang: [r["ten"] for r in server._luong_dong_tu_nhan_vien(hd, nv, 0, nam, thang)]
assert ten(2025, 3) == ["Còn làm", "Nghỉ 6/2025", "Nghỉ ngày đầy đủ"], "Tháng nghỉ việc vẫn còn trên bảng lương"
assert ten(2025, 4) == ["Còn làm", "Nghỉ 6/2025"]
assert ten(2025, 6) == ["Còn làm", "Nghỉ 6/2025"] and ten(2025, 7) == ["Còn làm"]
assert ten(2025, 1) == ["Còn làm", "Nghỉ 6/2025", "Nghỉ ngày đầy đủ"] and ten(2024, 12) == ["Còn làm", "Nghỉ 6/2025", "Nghỉ ngày đầy đủ", "Nghỉ 12/2024"]
assert len(server._luong_dong_tu_nhan_vien(hd, nv)) == 4, "Không có tháng -> lấy đủ (không lọc)"
assert server._luong_thang_nghi_viec("06/2025") == (2025, 6) and server._luong_thang_nghi_viec("") is None and server._luong_thang_nghi_viec("abc") is None
# vào làm vẫn theo quy tắc cũ (sau ngày 18 -> tháng sau)
assert server._luong_bat_dau_bhxh("20/03/2025") == (2025, 4) and server._luong_bat_dau_bhxh("10/2024") == (2024, 10)
# kế hoạch chi phí cả năm: người đã nghỉ không được chọn ở các tháng sau khi nghỉ
import random
hd2 = ["Mã NV", "Họ và tên", "Tháng/Năm nghỉ việc", "Lương Cơ bản", "Đóng BHXH"]
nv2 = [[str(i), f"NV{i}", "07/2025" if i < 4 else "", 5_310_000, "x"] for i in range(1, 8)]
pool = lambda t: server._luong_dong_tu_nhan_vien(hd2, nv2, 0, 2025, int(t))
th, tom = server._luong_ke_hoach(pool, 2025, 6, 9, 100_000_000, None, 50, 0, random.Random(1))
for t in ("08", "09"):
    assert all(r["ten"] not in ("NV1", "NV2", "NV3") for r in th[t]), "Đã nghỉ từ tháng 7 -> không xuất hiện ở tháng 8, 9"
assert sum(r["chi_phi_luong"] for rows in th.values() for r in rows) == 100_000_000
print("PASS 5: cột Tháng/Năm nghỉ việc: còn lên bảng lương tới hết tháng nghỉ, sau đó không còn (nạp danh sách + kế hoạch cả năm).")

# 6: import Excel Danh Sách Nhân Viên nhận cột "Ngày nghỉ việc" -> Tháng/Năm nghỉ việc
wb = openpyxl.Workbook()
w = wb.active
w.append(["Mã NV", "Họ và tên", "Ngày vào làm", "Ngày nghỉ việc", "Chức vụ", "Lương cơ bản"])
w.append([])
w.append(["1", "Nguyễn A", "01/01/2024", "30/06/2025", "KD", 5_310_000])
buf = io.BytesIO()
wb.save(buf)
kq = asyncio.run(server.nhap_lieu_import_nhan_vien(1, Req()))       # Up.read() đọc biến buf hiện tại
r = dict(zip(kq["header"], kq["rows"][0]))
assert r["Tháng/Năm nghỉ việc"] == "30/06/2025" and r["Tháng/Năm vào làm"] == "01/01/2024" and r["Chức vụ"] == "KD" and r["Lương Cơ bản"] == 5_310_000
print("PASS 6: import Excel nhận cột Ngày nghỉ việc, không nhầm sang cột khác.")

# 7: BHXH theo THỜI GIAN THAM GIA: tick + Tháng/Năm vào làm + Tháng/Năm nghỉ việc
hd7 = ["Mã NV", "Họ và tên", "Tháng/Năm vào làm", "Đóng BHXH", "Tháng/Năm nghỉ việc", "Lương Cơ bản"]
def bh(tick, vao, nghi, nam, thang):
    r = server._luong_dong_tu_nhan_vien(hd7, [["1", "A", vao, tick, nghi, 5_310_000]], 0, nam, thang)
    return r[0]["dong_bh"] if r else None          # None = không lên bảng lương tháng đó
assert [bh("x", "12/2024", "06/2025", 2024, m_) for m_ in (11, 12)] == [0, 1], "Trước tháng tham gia: chưa đóng; từ tháng vào làm: đóng"
assert [bh("x", "12/2024", "06/2025", 2025, m_) for m_ in (1, 6, 7, 8)] == [1, 1, None, None], "Đóng hết tháng nghỉ việc; sau đó không còn trên bảng lương"
assert bh("", "12/2024", "", 2025, 5) == 0 and bh("x", "12/2024", "", 2026, 3) == 1, "Không tick thì không đóng; không có ngày nghỉ = còn đóng"
assert bh("x", "", "", 2025, 1) == 1, "Không có ngày vào làm = đóng từ đầu"
# nghỉ việc ghi rõ ngày: < 14 ngày làm trong tháng nghỉ -> tháng nghỉ không đóng; >= 14 thì đóng
assert bh("x", "12/2024", "10/06/2025", 2025, 5) == 1 and bh("x", "12/2024", "10/06/2025", 2025, 6) == 0 and bh("x", "12/2024", "10/06/2025", 2025, 7) is None
assert bh("x", "12/2024", "20/06/2025", 2025, 6) == 1 and bh("x", "12/2024", "14/06/2025", 2025, 6) == 1 and bh("x", "12/2024", "13/06/2025", 2025, 6) == 0
assert bh("x", "12/2024", "06/2025", 2025, 6) == 1, "Chỉ ghi tháng/năm -> tính đủ tháng nghỉ"
assert server._luong_ngay_cu_the("06/2025") is None and server._luong_ngay_cu_the("20/06/2025") == 20 and server._luong_ngay_cu_the("2025-06-09") == 9
# vào làm sau ngày 18 vẫn tính từ tháng sau (quy tắc cũ), kết hợp nghỉ việc
assert [bh("x", "20/03/2025", "05/2025", 2025, m_) for m_ in (3, 4, 5, 6)] == [0, 1, 1, None]
# danh sách người đã nghỉ + API
hd8 = ["Mã NV", "Họ và tên", "Tháng/Năm nghỉ việc"]
nv8 = [["1", "A", ""], ["2", "B", "06/2025"], ["3", "C", "2025-08-15"]]
assert [x["ten"] for x in server._luong_nv_da_nghi(hd8, nv8, 2025, 7)] == ["B"] and [x["ten"] for x in server._luong_nv_da_nghi(hd8, nv8, 2025, 9)] == ["B", "C"]
assert server._luong_nv_da_nghi(hd8, nv8, 2025, 6) == [] and server._luong_nv_da_nghi(["Mã NV", "Họ và tên"], nv8, 2025, 9) == []
server.nhap_lieu_get = lambda cid, loai="nv": {"header": hd7, "rows": [["1", "A", "12/2024", "x", "06/2025", 5_310_000], ["2", "B", "12/2024", "x", "", 5_310_000]]}
kq = server.bang_luong_tu_nhan_vien(1, 2025, 8)
assert [r["ten"] for r in kq["rows"]] == ["B"] and kq["da_nghi"] == [{"ma": "1", "ten": "A"}] and kq["rows"][0]["dong_bh"] == 1
assert server.bang_luong_tu_nhan_vien(1, 2025, 0)["da_nghi"] == [] and len(server.bang_luong_tu_nhan_vien(1, 2025, 0)["rows"]) == 2
# kế hoạch cả năm: người BHXH chỉ được chọn đúng khoảng tham gia (vào 08/2025, nghỉ 10/2025)
hd9 = ["Mã NV", "Họ và tên", "Tháng/Năm vào làm", "Đóng BHXH", "Tháng/Năm nghỉ việc", "Lương Cơ bản"]
nv9 = [["1", "Dài hạn", "01/2024", "x", "", 5_310_000], ["2", "Ngắn hạn", "08/2025", "x", "10/2025", 5_310_000]]
pool9 = lambda t: server._luong_dong_tu_nhan_vien(hd9, nv9, 0, 2025, int(t))
th, tom = server._luong_ke_hoach(pool9, 2025, 7, 11, 100_000_000, None, 50, 0, random.Random(2))
for t, rows in th.items():
    ngan = [r for r in rows if r["ten"] == "Ngắn hạn"]
    if t == "11":
        assert not ngan, "Đã nghỉ từ tháng 11 -> không còn trên bảng lương"
    for r in ngan:
        if "08" <= t <= "10":
            assert r["dong_bh"] == 1 and r["bhxh_nld"] > 0, (t, r["dong_bh"])      # đúng khoảng tham gia: đóng BHXH
        else:
            assert r["dong_bh"] == 0 and r["bhxh_nld"] == 0, (t, r["dong_bh"])      # trước tháng tham gia: không đóng
assert any(r["ten"] == "Dài hạn" and r["dong_bh"] == 1 for rows in th.values() for r in rows)
print("PASS 7: BHXH theo thời gian tham gia: tick + vào làm + nghỉ việc (kể cả ngày nghỉ < 14), nạp/kế hoạch đúng theo tháng.")

print("\nALL DONE")
