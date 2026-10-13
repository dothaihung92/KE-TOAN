import json
import os
import sqlite3
import sys
import tempfile

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import server

# Xuất Excel theo kỳ: tờ khai nhập khẩu chỉ đưa vào "Chi tiết MUA VÀO" và "BK Mua vào" khi ngày đăng ký NẰM TRONG kỳ xuất
# (tu_ngay/den_ngay). Không truyền kỳ -> như cũ, đưa tất cả. (Lỗi thật: danh sách import chỉ hiện 3 tờ khai Quý 3 nhưng Excel xuất ra đủ mọi kỳ.)
tmp = tempfile.mkdtemp()
real = sqlite3.connect(os.path.join(tmp, "t.db"), check_same_thread=False)
real.row_factory = sqlite3.Row


class KD:
    def __init__(self, c): self.c = c
    def execute(self, *a): return self.c.execute(*a)
    def executemany(self, *a): return self.c.executemany(*a)
    def executescript(self, *a): return self.c.executescript(*a)
    def commit(self): self.c.commit()
    def close(self): pass
    def __getattr__(self, n): return getattr(self.c, n)


server.db = lambda: KD(real)
server.DOWNLOAD_DIR = os.path.join(tmp, "dl")
os.makedirs(server.DOWNLOAD_DIR, exist_ok=True)
server.init_db()
real.execute("INSERT INTO companies (id, ten, mst) VALUES (7,'CT','0315502878')")
for so, ng in (("TK_Q2", "2026-06-20"), ("TK_Q3a", "2026-08-07"), ("TK_Q3b", "2026-09-24"), ("TK_Q4", "2026-10-05")):
    real.execute("INSERT INTO tokhai_nhap (company_id, so_tk, ngay_dk, nguoi_xk, items_json, updated_at) VALUES (7,?,?,'XK',?,'')",
                 (so, ng, json.dumps([{"ten": "H", "tri_gia_gtgt": 1000, "tien_thue_gtgt": 100, "sluong": 1}])))
real.commit()


def so_tk_trong_file(tu, den):
    r = server.export_excel(7, 0, tu, den, 0, 0)
    wb = openpyxl.load_workbook(r.path, data_only=True)
    chi_tiet = {str(row[1]) for row in wb["Chi tiết MUA VÀO"].iter_rows(min_row=2, values_only=True) if row[0] == "TKNK"}
    bk = {str(row[2]) for row in wb["BK Mua vào"].iter_rows(min_row=2, values_only=True) if row[1] == "TKNK"}
    return chi_tiet, bk


ct, bk = so_tk_trong_file("01/07/2026", "30/09/2026")
assert ct == {"TK_Q3a", "TK_Q3b"} and bk == {"TK_Q3a", "TK_Q3b"}, (ct, bk)
print("PASS 1: xuất Excel Quý 3 chỉ có tờ khai nhập khẩu đăng ký trong quý 3 (cả Chi tiết MUA VÀO và BK Mua vào).")
ct, bk = so_tk_trong_file("01/01/2026", "31/12/2026")
assert ct == bk == {"TK_Q2", "TK_Q3a", "TK_Q3b", "TK_Q4"}, (ct, bk)
print("PASS 2: xuất cả năm có đủ các tờ khai trong năm.")
ct, bk = so_tk_trong_file("", "")
assert ct == bk == {"TK_Q2", "TK_Q3a", "TK_Q3b", "TK_Q4"}, (ct, bk)
print("PASS 3: không truyền kỳ -> như cũ, đưa tất cả tờ khai.")
print("ALL DONE")
