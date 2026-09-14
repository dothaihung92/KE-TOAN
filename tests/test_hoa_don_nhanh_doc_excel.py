import os
_REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
src = open(os.path.join(_REPO_ROOT, 'server.py'), encoding='utf-8').read()

# Regression test cho _doc_hoa_don_nhanh_tu_excel()/_hdn_so() trong server.py — tính năng mới theo
# yêu cầu người dùng: "có cách nào để phần mềm tự động hỗ trợ xuất hoá đơn... khách hàng sẽ nhắn
# thông tin hoặc gửi file excel hoặc pdf qua zalo" -> hỗ trợ soạn NHÁP hoá đơn từ file Excel khách
# gửi (dò tên/MST/địa chỉ + bảng hàng hoá theo heuristic, không có mẫu cố định vì mỗi khách gửi 1
# kiểu khác nhau), người dùng xem lại/sửa trên giao diện rồi tự phát hành thật.


def extract_fn(name):
    idx = src.index('def ' + name + '(')
    i = src.index(':', idx)
    lines = src[i + 1:].split('\n')
    body = []
    started = False
    for ln in lines:
        if ln.strip() == '' and not started:
            body.append(ln)
            continue
        if ln and not ln[0].isspace() and started:
            break
        if ln.strip():
            started = True
        body.append(ln)
    return src[idx:i + 1] + '\n'.join(body)


def extract_dict(name):
    idx = src.index(name + ' = {')
    depth = 0
    i = src.index('{', idx)
    j = i
    while True:
        if src[j] == '{':
            depth += 1
        elif src[j] == '}':
            depth -= 1
            if depth == 0:
                break
        j += 1
    return src[idx:j + 1]


ns = {}
exec(extract_dict('_HDN_LABELS_DAU'), ns)
exec(extract_dict('_HDN_LABELS_HANG'), ns)
exec(extract_fn('_hdn_so'), ns)
exec(extract_fn('_doc_hoa_don_nhanh_tu_excel'), ns)
_hdn_so = ns['_hdn_so']
_doc_hoa_don_nhanh_tu_excel = ns['_doc_hoa_don_nhanh_tu_excel']

import openpyxl

# ----- Test 1: layout "Nhãn: Giá trị" cùng 1 ô, bảng hàng hoá đủ 5 cột chuẩn. -----
wb = openpyxl.Workbook()
ws = wb.active
ws.append(["Tên công ty:", "CÔNG TY TNHH ABC XYZ"])
ws.append(["MST:", "0312345678"])
ws.append(["Địa chỉ:", "123 Nguyễn Huệ, Q1, TP.HCM"])
ws.append([])
ws.append(["Tên hàng hóa", "ĐVT", "Số lượng", "Đơn giá", "Thành tiền"])
ws.append(["Bàn ghế văn phòng", "Bộ", 5, 1750000, 8750000])
ws.append(["Tủ hồ sơ sắt", "Cái", 3, 1550000, 4650000])
d = _doc_hoa_don_nhanh_tu_excel(wb)
assert d["ten"] == "CÔNG TY TNHH ABC XYZ", d["ten"]
assert d["mst"] == "0312345678", d["mst"]
assert d["diaChi"] == "123 Nguyễn Huệ, Q1, TP.HCM", d["diaChi"]
assert len(d["items"]) == 2, d["items"]
assert d["items"][0] == {"ten": "Bàn ghế văn phòng", "dvt": "Bộ", "sl": 5.0, "dgia": 1750000.0, "thanhTien": 8750000.0}, d["items"][0]
assert d["items"][1]["thanhTien"] == 4650000.0
print("PASS 1: layout 'Nhãn: Giá trị' cùng ô + bảng hàng hoá đủ cột -> dò đúng tên/MST/địa chỉ/hàng hoá.")

# ----- Test 2: layout nhãn/giá trị TÁCH 2 ô riêng (không có dấu ':'), thiếu cột "Thành tiền" ->
# phải TỰ TÍNH thanhTien = sl * dgia. -----
wb2 = openpyxl.Workbook()
ws2 = wb2.active
ws2.append(["Khách hàng", "CÔNG TY CP DEMO"])
ws2.append(["Mã số thuế", "0399998888"])
ws2.append(["Đ/c", "45 Lê Lợi, Q1"])
ws2.append([])
ws2.append(["STT", "Sản phẩm", "Đơn vị", "SL", "Đơn giá"])
ws2.append([1, "Giấy A4", "Ram", 200, 60000])
ws2.append([2, "Mực in", "Hộp", 15, 300000])
d2 = _doc_hoa_don_nhanh_tu_excel(wb2)
assert d2["ten"] == "CÔNG TY CP DEMO", d2["ten"]
assert d2["mst"] == "0399998888", d2["mst"]
assert d2["diaChi"] == "45 Lê Lợi, Q1", d2["diaChi"]
assert len(d2["items"]) == 2
assert d2["items"][0]["thanhTien"] == 12000000.0, d2["items"][0]  # 200 * 60000, tu tinh vi thieu cot
assert d2["items"][1]["thanhTien"] == 4500000.0
print("PASS 2: nhãn/giá trị tách 2 ô + thiếu cột 'Thành tiền' -> vẫn dò đúng, tự tính thanhTien = SL x Đơn giá.")

# ----- Test 3 (không hồi quy/an toàn dữ liệu lạ): sheet trống hoàn toàn -> không crash, trả rỗng. -----
wb3 = openpyxl.Workbook()
d3 = _doc_hoa_don_nhanh_tu_excel(wb3)
assert d3["ten"] == "" and d3["mst"] == "" and d3["diaChi"] == "" and d3["items"] == [], d3
print("PASS 3: sheet trống -> không crash, trả về rỗng an toàn.")

# ----- Test 4: _hdn_so() đọc đúng nhiều định dạng số khác nhau. -----
assert _hdn_so(1750000) == 1750000.0
assert _hdn_so("1750000") == 1750000.0
assert _hdn_so("1.750.000") == 1750000.0  # dấu . la phan cach nghin kieu VN
assert _hdn_so("1,750,000") == 1750000.0  # dau , la phan cach nghin kieu US
assert _hdn_so("1.750.000,50") == 1750000.50  # VN: . nghin, , thap phan
assert _hdn_so("1,750,000.50") == 1750000.50  # US: , nghin, . thap phan
assert _hdn_so("") == 0.0
assert _hdn_so(None) == 0.0
assert _hdn_so("abc") == 0.0  # khong doan duoc -> 0, khong crash
print("PASS 4: _hdn_so() đọc đúng nhiều định dạng số (VN/US, number thật, rỗng, chuỗi lạ không crash).")

print("\nALL DONE")
